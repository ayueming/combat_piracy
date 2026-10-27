# page_links.py
import sqlite3
import os
import subprocess
import json
from PySide6.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
                               QHeaderView, QMessageBox, QButtonGroup, QAbstractItemView,
                               QFileDialog, QTableView, QDialog, QGroupBox,
                               QStyledItemDelegate, QStyle, QScrollArea, QFrame)
from PySide6.QtCore import Qt, QSettings, QItemSelectionModel, QEvent, QTimer
from PySide6.QtGui import QColor, QFont, QStandardItemModel, QStandardItem, QPainter
from db_config import get_conn
from feature_extractor import extract_pure_feature, get_platform_category
from data_importer import import_sites_from_file

from ui_components import (
    ExcelFilterProxyModel,
    FilterHeaderView,
    CenteredCheckBoxDelegate,
    ComboCellDelegate,
    StatusComboDelegate,
    HintButton,
)
from complaint_service import prepare_complaints_data

from app_data import PLATFORM_CONFIG, ALL_SUB_PLATFORMS, CAT_TO_CODE_MAP
from app_data.complaint_platform_consts import HIDDEN_MAIN_CATEGORIES
from app_data.crawl_status_consts import CRAWL_STATUS_DONE

from ui_styles import (
    BTN_SECONDARY,
    BTN_SECONDARY_DARK,
    BTN_PRIMARY,
    BTN_DANGER,
    BTN_DANGER_OUTLINE,
    BTN_TOGGLE,
    BTN_TOGGLE_MAIN,
)

from app_signals import signals

# ★ 列索引
COL_CHK = 0
COL_BOOK = 1
COL_CAT = 2
COL_URL = 3
COL_PUBLISHER = 4  # 网站域名（发布者，只读）
COL_ATTR = 5  # 属性（手动获取/机器人获取）
COL_CODE = 6
COL_CRAWL_STATUS = 7
COL_COMP_TIME = 8
COL_RES_STATUS = 9
COL_TITLE = 10  # ★ 新增：标题（只读）

TABLE_HEADERS = [
    "勾选", "关联作品(必填)", "资源分类(必填)", "资源链接(必填)", "网站域名",
    "属性", "特征码", "抓取状态", "投诉时间", "资源状态", "标题",
]


class LinksTableView(QTableView):
    def mousePressEvent(self, event):
        index = self.indexAt(event.pos())
        if index.isValid():
            if index.column() == COL_CHK:
                sel_model = self.selectionModel()
                is_selected = sel_model.isRowSelected(index.row(), index.parent())
                mods = event.modifiers()

                self.setCurrentIndex(index)

                if mods & (Qt.ControlModifier | Qt.ShiftModifier):
                    if is_selected:
                        sel_model.select(index, QItemSelectionModel.Deselect | QItemSelectionModel.Rows)
                    else:
                        sel_model.select(index, QItemSelectionModel.Select | QItemSelectionModel.Rows)
                else:
                    if is_selected:
                        sel_model.select(index, QItemSelectionModel.Deselect | QItemSelectionModel.Rows)
                    else:
                        sel_model.select(index, QItemSelectionModel.ClearAndSelect | QItemSelectionModel.Rows)

                event.accept()
                return

            elif index.column() in (COL_BOOK, COL_CAT, COL_RES_STATUS):
                model = self.model()
                if index.column() == COL_RES_STATUS:
                    comp_time = model.index(index.row(), COL_COMP_TIME).data()
                    if comp_time and str(comp_time).strip():
                        return

                if self.selectionModel().isRowSelected(index.row(), index.parent()):
                    self.edit(index)
                else:
                    self.setCurrentIndex(index)
                event.accept()
                return

        super().mousePressEvent(event)


class ImportLinksDialog(QDialog):
    def __init__(self, parent_page):
        super().__init__(parent_page)
        self.parent_page = parent_page
        self.setWindowTitle("批量导入盗版链接")
        self.setFixedSize(480, 300)
        self.setup_ui()

    def setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(15)

        group_template = QGroupBox("第一步：下载标准模板")
        group_template.setStyleSheet("QGroupBox { font-weight: bold; color: #333; }")
        layout_template = QVBoxLayout(group_template)

        lbl_tip1 = QLabel("为了确保数据能正确导入，请先下载模板，并按照表头格式填写您的链接信息。")
        lbl_tip1.setStyleSheet("color: #666; font-weight: normal;")
        lbl_tip1.setWordWrap(True)
        layout_template.addWidget(lbl_tip1)

        btn_download = QPushButton("下载 Excel 模板")
        btn_download.setStyleSheet(BTN_SECONDARY_DARK + "QPushButton { padding: 8px 12px; }")
        btn_download.clicked.connect(self.download_template)
        layout_template.addWidget(btn_download)

        layout.addWidget(group_template)

        group_upload = QGroupBox("第二步：上传已填写的文件")
        group_upload.setStyleSheet("QGroupBox { font-weight: bold; color: #333; }")
        layout_upload = QVBoxLayout(group_upload)

        lbl_tip2 = QLabel("请选择填写好的 Excel (.xlsx) 或 CSV (.csv) 文件进行导入。如果特征码已存在，系统会自动跳过。")
        lbl_tip2.setStyleSheet("color: #666; font-weight: normal;")
        lbl_tip2.setWordWrap(True)
        layout_upload.addWidget(lbl_tip2)

        btn_upload = QPushButton("选择文件并导入")
        btn_upload.setStyleSheet(BTN_PRIMARY + "QPushButton { padding: 8px 12px; }")
        btn_upload.clicked.connect(self.upload_file)
        layout_upload.addWidget(btn_upload)

        layout.addWidget(group_upload)
        layout.addStretch()

    def download_template(self):
        path, _ = QFileDialog.getSaveFileName(self, "保存链接导入模板", "链接导入模板.xlsx", "Excel 文件 (*.xlsx)")
        if not path: return
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "盗版链接"

            headers = ["关联作品(必填)", "资源分类", "资源链接(必填)", "网站域名", "属性", "投诉时间", "状态"]
            for col, header in enumerate(headers, 1):
                ws.cell(row=1, column=col, value=header)

            for col in range(1, 8):
                ws.column_dimensions[chr(64 + col)].width = 20

            wb.save(path)
            QMessageBox.information(self, "下载成功", f"标准模板已保存至：\n{path}")
        except Exception as e:
            QMessageBox.critical(self, "失败", f"保存模板失败：{e}")

    def upload_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "选择导入文件", "",
                                              "表格文件 (*.csv *.xlsx);;CSV 文件 (*.csv);;Excel 文件 (*.xlsx)")
        if not path: return
        try:
            success, duplicate, errors = import_sites_from_file(path)
            self.parent_page.refresh_table()

            msg = f"成功导入了 {success} 条新记录！\n(自动跳过了 {duplicate} 条重复)"
            if errors:
                msg += f"\n\n⚠️ 有 {len(errors)} 条记录因校验失败被跳过：\n\n"
                preview = errors[:10]
                msg += "\n".join(preview)
                if len(errors) > 10:
                    msg += f"\n...（还有 {len(errors) - 10} 条未展示）"

            QMessageBox.information(self, "导入完成", msg)
            self.accept()
        except Exception as e:
            QMessageBox.critical(self, "导入失败", str(e))


class LinksPage(QWidget):
    def __init__(self, main_window):
        super().__init__()
        self.main_window = main_window
        self.settings = QSettings("AuthorTools", "AntiPiracy")

        self.cat_map = {"全部链接": ["全部资源"]}
        for main_cat, sub_list in PLATFORM_CONFIG.items():
            sub_names = [sub[0] for sub in sub_list]
            self.cat_map[main_cat] = sub_names
            for sub_name in sub_names:
                self.cat_map[sub_name] = [sub_name]

        self.book_list = []
        self._is_updating = False
        self._current_cat = "全部链接"
        self.filters_by_cat = {}

        self.setup_ui()

        signals.link_saved.connect(self.on_link_saved_signal, Qt.UniqueConnection)

    def showEvent(self, event):
        super().showEvent(event)
        self.refresh_books()
        self.refresh_table()

    def on_link_saved_signal(self, record, duplicate):
        self.append_site_row(record, duplicate)

    def setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(10)  # 稍微缩小一点间距，让整体更紧凑

        self.cat_group = QButtonGroup(self)
        self.cat_group.setExclusive(True)

        filter_v_layout = QVBoxLayout()
        filter_v_layout.setSpacing(8)  # 缩小间距

        # ------------------ 快捷分组 ------------------
        main_row = QHBoxLayout()
        main_row.setSpacing(8)
        main_row.setContentsMargins(0, 0, 0, 0)

        # 1. 左侧固定的标签
        lbl_main = QLabel("快捷分组:")
        lbl_main.setStyleSheet("color: #666; font-weight: bold; font-size: 13px;")
        lbl_main.setFixedWidth(65)
        main_row.addWidget(lbl_main)

        # 2. 右侧的滑动区域
        main_scroll = QScrollArea()
        main_scroll.setWidgetResizable(True)
        main_scroll.setFrameShape(QFrame.NoFrame)
        main_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        main_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        main_scroll.setFixedHeight(45)  # 原来 60 太高了，改小点去掉多余灰底
        # ★ 核心：去除 ScrollArea 和其内部的灰色底纹
        main_scroll.setStyleSheet("QScrollArea { border: none; background: transparent; }")
        main_scroll.viewport().setStyleSheet("background: transparent;")

        main_widget = QWidget()
        main_widget.setStyleSheet("background: transparent;")  # ★ 容器透明
        btn_layout_main = QHBoxLayout(main_widget)
        btn_layout_main.setContentsMargins(0, 0, 0, 0)
        btn_layout_main.setSpacing(8)

        # 添加按钮到内部布局
        all_btn = QPushButton("全部链接")
        all_btn.setCheckable(True)
        all_btn.setStyleSheet(BTN_TOGGLE_MAIN)
        all_btn.setCursor(Qt.PointingHandCursor)
        all_btn.setFixedSize(85, 32)
        self.cat_group.addButton(all_btn)
        btn_layout_main.addWidget(all_btn)
        all_btn.setChecked(True)

        for main_cat in PLATFORM_CONFIG.keys():
            if main_cat in HIDDEN_MAIN_CATEGORIES:
                continue
            btn = QPushButton(main_cat)
            btn.setCheckable(True)
            btn.setStyleSheet(BTN_TOGGLE_MAIN)
            btn.setCursor(Qt.PointingHandCursor)
            btn.setFixedSize(85, 32)
            self.cat_group.addButton(btn)
            btn_layout_main.addWidget(btn)

        btn_layout_main.addStretch()

        main_scroll.setWidget(main_widget)
        main_row.addWidget(main_scroll, 1)  # 滚动区占满剩余空间

        filter_v_layout.addLayout(main_row)

        # ------------------ 具体平台 ------------------
        sub_row = QHBoxLayout()
        sub_row.setSpacing(8)
        sub_row.setContentsMargins(0, 0, 0, 0)

        # 1. 左侧固定的标签
        lbl_sub = QLabel("具体平台:")
        lbl_sub.setStyleSheet("color: #666; font-weight: bold; font-size: 13px;")
        lbl_sub.setFixedWidth(65)
        sub_row.addWidget(lbl_sub)

        # 2. 右侧的滑动区域
        sub_scroll = QScrollArea()
        sub_scroll.setWidgetResizable(True)
        sub_scroll.setFrameShape(QFrame.NoFrame)
        sub_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        sub_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        sub_scroll.setFixedHeight(45)
        sub_scroll.setStyleSheet("QScrollArea { border: none; background: transparent; }")
        sub_scroll.viewport().setStyleSheet("background: transparent;")

        sub_widget = QWidget()
        sub_widget.setStyleSheet("background: transparent;")
        btn_layout_sub = QHBoxLayout(sub_widget)
        btn_layout_sub.setContentsMargins(0, 0, 0, 0)
        btn_layout_sub.setSpacing(8)

        for sub_name in ALL_SUB_PLATFORMS:
            if sub_name in HIDDEN_MAIN_CATEGORIES:
                continue
            btn = QPushButton(sub_name)
            btn.setCheckable(True)
            btn.setStyleSheet(BTN_TOGGLE)
            btn.setCursor(Qt.PointingHandCursor)
            btn.setFixedSize(85, 32)
            self.cat_group.addButton(btn)
            btn_layout_sub.addWidget(btn)

        btn_layout_sub.addStretch()

        sub_scroll.setWidget(sub_widget)
        sub_row.addWidget(sub_scroll, 1)  # 滚动区占满剩余空间

        filter_v_layout.addLayout(sub_row)

        layout.addLayout(filter_v_layout)

        # ------------------ 操作按钮 ------------------
        action_layout = QHBoxLayout()
        action_layout.setContentsMargins(0, 0, 0, 0)
        action_layout.addWidget(HintButton("带有淡红底色的单元格为必填项未填写。点击资源状态列可展开修改。"))

        action_layout.addStretch()

        self.stats_label = QLabel("")
        self.stats_label.setStyleSheet("color: #666; font-size: 13px;")
        action_layout.addWidget(self.stats_label)

        self.report_btn = QPushButton(" 一键发起投诉 (勾选)")
        self.report_btn.setStyleSheet(BTN_DANGER)
        self.report_btn.clicked.connect(self.send_to_complaints)
        action_layout.addWidget(self.report_btn)

        import_btn = QPushButton("批量导入表格")
        import_btn.setStyleSheet(BTN_SECONDARY)
        import_btn.clicked.connect(self.open_import_dialog)
        action_layout.addWidget(import_btn)

        export_btn = QPushButton("导出勾选数据")
        export_btn.setStyleSheet(BTN_SECONDARY)
        export_btn.clicked.connect(self.export_to_excel)
        action_layout.addWidget(export_btn)

        del_btn = QPushButton("🗑️ 删除选中")
        del_btn.setStyleSheet(BTN_DANGER_OUTLINE)
        del_btn.clicked.connect(self.delete_site)
        action_layout.addWidget(del_btn)

        # ★ 用 QScrollArea 包裹操作按钮
        action_widget = QWidget()
        action_widget.setStyleSheet("background: transparent;")  # ★ 让内部容器透明
        action_widget.setLayout(action_layout)

        action_scroll = QScrollArea()
        action_scroll.setWidgetResizable(True)
        action_scroll.setFrameShape(QFrame.NoFrame)
        action_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        action_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        action_scroll.setFixedHeight(50)  # 略微调小一点，刚刚好放下按钮即可
        action_scroll.setStyleSheet("QScrollArea { border: none; background: transparent; }")
        action_scroll.viewport().setStyleSheet("background: transparent;")  # ★ 核心：让滚动视图区透明
        action_scroll.setWidget(action_widget)
        layout.addWidget(action_scroll)

        # ------------------ 表格 ------------------
        self.model = QStandardItemModel(0, len(TABLE_HEADERS), self)
        self.model.setHorizontalHeaderLabels(TABLE_HEADERS)
        self.model.itemChanged.connect(self.on_item_changed)

        self.proxy_model = ExcelFilterProxyModel(self)
        self.proxy_model.setSourceModel(self.model)
        self.proxy_model.filterChanged.connect(self._on_filter_changed)

        self.table = LinksTableView()
        self.table.setModel(self.proxy_model)

        custom_header = FilterHeaderView(Qt.Horizontal, self.table)
        self.table.setHorizontalHeader(custom_header)

        header = self.table.horizontalHeader()
        header.setFixedHeight(35)
        header.setSectionResizeMode(QHeaderView.Interactive)

        self.table.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOn)
        self.table.verticalHeader().setFixedWidth(35)

        self.table.setColumnWidth(COL_CHK, 40)
        self.table.setColumnWidth(COL_BOOK, 120)
        self.table.setColumnWidth(COL_CAT, 100)
        self.table.setColumnWidth(COL_URL, 400)
        self.table.setColumnWidth(COL_PUBLISHER, 120)
        self.table.setColumnWidth(COL_ATTR, 100)
        self.table.setColumnWidth(COL_CODE, 120)
        self.table.setColumnWidth(COL_CRAWL_STATUS, 100)
        self.table.setColumnWidth(COL_COMP_TIME, 120)
        self.table.setColumnWidth(COL_RES_STATUS, 100)
        self.table.setColumnWidth(COL_TITLE, 250)

        self.table.setItemDelegateForColumn(COL_CHK, CenteredCheckBoxDelegate(self.table))
        self.table.setItemDelegateForColumn(COL_BOOK, ComboCellDelegate(
            options=lambda: [(name, name) for _, name, _, _ in self.book_list],
            parent=self.table,
        ))
        self.table.setItemDelegateForColumn(COL_CAT, ComboCellDelegate(
            options=[(c, c) for c in ALL_SUB_PLATFORMS],
            parent=self.table,
        ))
        self.table.setItemDelegateForColumn(COL_RES_STATUS, StatusComboDelegate(self.table))

        self.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setEditTriggers(
            QAbstractItemView.DoubleClicked | QAbstractItemView.SelectedClicked | QAbstractItemView.EditKeyPressed)
        self.table.setMouseTracking(True)
        self.table.selectionModel().selectionChanged.connect(self.on_selection_changed)
        layout.addWidget(self.table, 1)

        self.table.horizontalHeader().sectionResized.connect(self._save_column_widths)
        self._load_column_widths()

        self.cat_group.buttonClicked.connect(self._on_cat_clicked)
        self._load_ui_state()

        self.refresh_books()
        self.refresh_table()

    def _save_column_widths(self):
        if not hasattr(self, '_col_width_timer'):
            self._col_width_timer = QTimer(self)
            self._col_width_timer.setSingleShot(True)
            self._col_width_timer.timeout.connect(self._do_save_column_widths)
        self._col_width_timer.start(400)

    def _do_save_column_widths(self):
        header = self.table.horizontalHeader()
        widths = {}
        for col in range(self.model.columnCount()):
            if col == COL_URL: continue
            widths[str(col)] = header.sectionSize(col)
        self.settings.setValue("links_col_widths", json.dumps(widths))

    def _load_column_widths(self):
        raw = self.settings.value("links_col_widths", "")
        if raw:
            try:
                widths = json.loads(raw)
                header = self.table.horizontalHeader()
                header.blockSignals(True)
                for col_str, width in widths.items():
                    col = int(col_str)
                    if col != COL_URL:
                        header.resizeSection(col, width)
                header.blockSignals(False)
            except Exception:
                pass

    def _load_ui_state(self):
        self._current_cat = self.settings.value("links_current_cat", "全部链接")
        for btn in self.cat_group.buttons():
            if btn.text() == self._current_cat:
                btn.setChecked(True)
                break
        else:
            first_btn = self.cat_group.buttons()[0]
            first_btn.setChecked(True)
            self._current_cat = first_btn.text()
            self.settings.setValue("links_current_cat", self._current_cat)

        raw = self.settings.value("links_filters_by_cat", "")
        if raw:
            try:
                self.filters_by_cat = json.loads(raw)
            except Exception:
                self.filters_by_cat = {}

        self._apply_filters_to_proxy()

    def _on_cat_clicked(self, btn):
        new_cat = btn.text()
        if new_cat == self._current_cat:
            return

        self._current_cat = new_cat
        self._apply_filters_to_proxy()
        self.settings.setValue("links_current_cat", self._current_cat)
        self.refresh_table()

    def _apply_filters_to_proxy(self):
        self._is_updating = True
        self.proxy_model.filter_dict.clear()

        cat_filter = self.filters_by_cat.get(self._current_cat, {})
        for col_str, allowed_list in cat_filter.items():
            self.proxy_model.filter_dict[int(col_str)] = set(allowed_list)

        self.proxy_model.invalidateFilter()
        self.table.horizontalHeader().viewport().update()
        self._is_updating = False

    def _on_filter_changed(self):
        if self._is_updating:
            return

        current_filter = {}
        for col, allowed_set in self.proxy_model.filter_dict.items():
            current_filter[str(col)] = list(allowed_set)

        self.filters_by_cat[self._current_cat] = current_filter
        self.settings.setValue("links_filters_by_cat", json.dumps(self.filters_by_cat))
        self._update_stats()

    def _update_stats(self):
        total_count = max(0, self.model.rowCount() - 15)
        visible_count = max(0, self.proxy_model.rowCount() - 15)
        has_filter = bool(self.proxy_model.filter_dict)

        if not has_filter:
            self.stats_label.setText(f"共 {total_count} 条链接")
            self.stats_label.setStyleSheet("color: #666; font-size: 13px;")
        elif visible_count == 0:
            self.stats_label.setText(f"⚠️ 当前筛选无结果 (共 {total_count} 条)")
            self.stats_label.setStyleSheet("color: #E74C3C; font-size: 13px; font-weight: bold;")
        else:
            self.stats_label.setText(f"筛选结果: {visible_count} / {total_count} 条")
            self.stats_label.setStyleSheet("color: #10B981; font-size: 13px; font-weight: bold;")

    def open_import_dialog(self):
        dialog = ImportLinksDialog(self)
        dialog.exec()

    def export_to_excel(self):
        checked_rows = set()
        for row in range(self.model.rowCount()):
            chk_item = self.model.item(row, COL_CHK)
            if chk_item and chk_item.checkState() == Qt.Checked:
                url_item = self.model.item(row, COL_URL)
                if url_item and url_item.text().strip():
                    checked_rows.add(row)

        for proxy_idx in self.table.selectionModel().selectedRows():
            src_idx = self.proxy_model.mapToSource(proxy_idx)
            if src_idx.isValid():
                url_item = self.model.item(src_idx.row(), COL_URL)
                if url_item and url_item.text().strip():
                    checked_rows.add(src_idx.row())

        if not checked_rows:
            QMessageBox.warning(self, "提示", "请先勾选或选中需要导出的数据！")
            return

        current_display = self._current_cat
        default_name = f"盗版链接导出_{current_display}.xlsx"
        path, _ = QFileDialog.getSaveFileName(self, "保存Excel文件", default_name, "Excel 文件 (*.xlsx)")
        if not path: return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment

            wb = Workbook()
            ws = wb.active
            ws.title = "盗版链接"

            headers = ["关联作品(必填)", "资源分类", "资源链接(必填)", "网站域名", "属性", "投诉时间", "状态"]
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")

            for row_idx, r in enumerate(sorted(list(checked_rows)), start=2):
                book_name = self.model.item(r, COL_BOOK).text().strip()
                cat = self.model.item(r, COL_CAT).text().strip()
                url = self.model.item(r, COL_URL).text().replace(" (重复资源链接)", "").strip()
                publisher = self.model.item(r, COL_PUBLISHER).text().strip()
                attr = self.model.item(r, COL_ATTR).text().strip()
                comp_time = self.model.item(r, COL_COMP_TIME).text().strip()
                res_status = self.model.item(r, COL_RES_STATUS).text().strip()

                ws.cell(row=row_idx, column=1, value=book_name)
                ws.cell(row=row_idx, column=2, value=cat)
                ws.cell(row=row_idx, column=3, value=url)
                ws.cell(row=row_idx, column=4, value=publisher)
                ws.cell(row=row_idx, column=5, value=attr)
                ws.cell(row=row_idx, column=6, value=comp_time)
                ws.cell(row=row_idx, column=7, value=res_status)

            ws.column_dimensions['A'].width = 20
            ws.column_dimensions['B'].width = 15
            ws.column_dimensions['C'].width = 50
            ws.column_dimensions['D'].width = 20
            ws.column_dimensions['E'].width = 20
            ws.column_dimensions['F'].width = 20
            ws.column_dimensions['G'].width = 15

            wb.save(path)
            QMessageBox.information(self, "导出成功", f"已成功导出 {len(checked_rows)} 条勾选数据！")
            folder_path = os.path.dirname(path)
            subprocess.Popen(f'explorer "{folder_path}"')
        except Exception as e:
            QMessageBox.critical(self, "导出失败", f"生成 Excel 时出错：{e}")

    def on_selection_changed(self, selected, deselected):
        if self._is_updating: return

        selected_rows = set(idx.row() for idx in selected.indexes())
        deselected_rows = set(idx.row() for idx in deselected.indexes())
        deselected_rows -= selected_rows

        self._is_updating = True
        try:
            for r in deselected_rows:
                src_idx = self.proxy_model.mapToSource(self.proxy_model.index(r, COL_CHK))
                item = self.model.item(src_idx.row(), COL_CHK)
                if item and item.data(Qt.UserRole) is not None:
                    if item.checkState() != Qt.Unchecked:
                        item.setCheckState(Qt.Unchecked)

            for r in selected_rows:
                src_idx = self.proxy_model.mapToSource(self.proxy_model.index(r, COL_CHK))
                item = self.model.item(src_idx.row(), COL_CHK)
                if item and item.data(Qt.UserRole) is not None:
                    if item.checkState() != Qt.Checked:
                        item.setCheckState(Qt.Checked)
        finally:
            self._is_updating = False

        self.table.horizontalHeader().viewport().update()

    def send_to_complaints(self):
        target_rows = set()
        for row in range(self.model.rowCount()):
            chk_item = self.model.item(row, COL_CHK)
            if chk_item and chk_item.checkState() == Qt.Checked:
                target_rows.add(row)

        for proxy_idx in self.table.selectionModel().selectedRows():
            src_idx = self.proxy_model.mapToSource(proxy_idx)
            if src_idx.isValid():
                target_rows.add(src_idx.row())

        if not target_rows:
            QMessageBox.warning(self, "提示", "请先勾选或选中需要投诉的条目！")
            return

        missing_rows = []
        for r in target_rows:
            book = self.model.item(r, COL_BOOK).text().strip()
            cat = self.model.item(r, COL_CAT).text().strip()
            url = self.model.item(r, COL_URL).text().replace(" (重复资源链接)", "").strip()

            if not book or not cat or not url:
                missing_rows.append(str(r + 1))

        if missing_rows:
            rows_text = "、".join(missing_rows[:10])
            if len(missing_rows) > 10:
                rows_text += " 等"
            QMessageBox.critical(
                self, "⛔ 拦截：必填项缺失",
                f"检测到以下行存在必填项未填写：\n\n👉 第 {rows_text} 行\n\n"
                f"请把红底单元格补充完整后再发起投诉！"
            )
            return

        batches, missing_books = prepare_complaints_data(self.model, self.book_list)

        if not batches:
            QMessageBox.warning(self, "提示", "未能生成投诉数据，请检查勾选情况！")
            return

        if missing_books:
            missing_str = "、".join(missing_books)
            QMessageBox.critical(
                self, "⛔ 拦截：作品信息不完整",
                f"检测到以下作品未配置【平台作品ID】或未上传【授权书PDF】或未填写【权利人】：\n\n"
                f"👉 {missing_str}\n\n"
                f"请先前往【📚 作品管理】页面，补充完整后再来发起投诉！\n"
            )
            return

        self.main_window.sidebar.setCurrentRow(4)
        self.main_window.complaint_page.load_batches(batches)

    def refresh_books(self):
        conn = get_conn()
        self.book_list = conn.execute("SELECT id, name, platform_id, evidence_path FROM books").fetchall()
        conn.close()

    def on_item_changed(self, item):
        if self._is_updating: return

        row = item.row()
        col = item.column()

        self.update_row_style(row)

        if col == COL_CHK:
            self.table.horizontalHeader().viewport().update()
            return

        if col == COL_URL:
            raw_url = item.text().strip()

            current_display = self._current_cat
            auto_cat = get_platform_category(raw_url)

            strict_cats = [name for name, info in CAT_TO_CODE_MAP.items() if info["module"] == "netdisk"]
            strict_cats.extend(["百度贴吧", "百度知道"])

            if current_display in ALL_SUB_PLATFORMS:
                if current_display in strict_cats and auto_cat != current_display:
                    msg = f"系统识别该链接可能不属于【{current_display}】。"
                    if auto_cat:
                        msg += f"（系统识别为【{auto_cat}】）"

                    reply = QMessageBox.question(
                        self, "⚠️ 分类不匹配提示",
                        f"{msg}\n\n请核对链接，是否强制将其保留在当前分类下？",
                        QMessageBox.Yes | QMessageBox.No, QMessageBox.No
                    )

                    if reply != QMessageBox.Yes:
                        self._is_updating = True
                        item.setText("")
                        self.model.item(row, COL_CAT).setText("")
                        self._is_updating = False
                        self.update_row_style(row)
                        self.save_row_to_db(row)
                        return

                self._is_updating = True
                self.model.item(row, COL_CAT).setText(current_display)

            elif current_display == "其他":
                self._is_updating = True
                self.model.item(row, COL_CAT).setText("其他")
            else:
                self._is_updating = True
                if auto_cat and not self.model.item(row, COL_CAT).text().strip():
                    self.model.item(row, COL_CAT).setText(auto_cat)

            self._is_updating = True
            pure_code = extract_pure_feature(raw_url)
            self.model.item(row, COL_CODE).setText(pure_code)

            db_id = self.model.item(row, COL_CHK).data(Qt.UserRole)
            conn = get_conn()
            cursor = conn.cursor()
            cursor.execute("SELECT COUNT(*) FROM sites WHERE feature_code = ? AND id != ?", (pure_code, db_id or "0"))
            is_duplicate = cursor.fetchone()[0] > 0
            conn.close()

            if is_duplicate:
                base_url = raw_url.replace(" (重复资源链接)", "")
                item.setText(f"{base_url} (重复资源链接)")
                item.setForeground(QColor("red"))
                item.setFont(QFont("Microsoft YaHei", 9, QFont.Bold))
            else:
                item.setForeground(QColor("#000000"))
                item.setFont(QFont("Microsoft YaHei", 9, QFont.Normal))
            self._is_updating = False

        self.save_row_to_db(row)

    def update_row_style(self, row):
        items = [self.model.item(row, c) for c in [COL_BOOK, COL_CAT, COL_URL]]
        if not all(items): return

        texts = [it.text().replace(" (重复资源链接)", "").strip() for it in items]
        db_id = self.model.item(row, COL_CHK).data(Qt.UserRole)

        is_blank_new_row = all(not t for t in texts) and not db_id

        bg_color = QColor("#FADBD8")
        transparent = QColor(0, 0, 0, 0)

        self._is_updating = True
        for i, item in enumerate(items):
            if not is_blank_new_row and not texts[i]:
                item.setBackground(bg_color)
            else:
                item.setBackground(transparent)
        self._is_updating = False

    def save_row_to_db(self, row):
        self._is_updating = True
        db_id = self.model.item(row, COL_CHK).data(Qt.UserRole)
        book_name = self.model.item(row, COL_BOOK).text().strip()
        cat = self.model.item(row, COL_CAT).text().strip()
        url = self.model.item(row, COL_URL).text().replace(" (重复资源链接)", "").strip()
        name = self.model.item(row, COL_TITLE).text().strip()  # ★ 标题从第 11 列读
        attr = self.model.item(row, COL_ATTR).text().strip()

        code = self.model.item(row, COL_CODE).text().strip()
        crawl_status = self.model.item(row, COL_CRAWL_STATUS).text().strip()
        comp_time = self.model.item(row, COL_COMP_TIME).text().strip()
        display_status = self.model.item(row, COL_RES_STATUS).text().strip()

        res_status = "已失效" if display_status == "已失效" else ""

        if not url and not book_name and not cat and not name and not attr and not comp_time:
            if db_id:
                conn = get_conn()
                conn.execute("DELETE FROM sites WHERE id=?", (db_id,))
                conn.commit()
                conn.close()

                self.model.item(row, COL_CHK).setData(None, Qt.UserRole)
                self.model.item(row, COL_CODE).setText("")
                self.update_row_style(row)

            self._is_updating = False
            return

        book_id = None
        if book_name:
            for b_id, b_n, _, _ in self.book_list:
                if b_n == book_name:
                    book_id = b_id
                    break

        conn = get_conn()
        cursor = conn.cursor()
        try:
            if not db_id:
                cursor.execute('''INSERT INTO sites
                                  (category, target_platform, site_url, name, attr, feature_code, complaint_time,
                                   resource_status,
                                   book_id, book_name, crawl_status)
                                  VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                               (cat, cat, url, name, attr, code, comp_time, res_status, book_id, book_name,
                                crawl_status))
                conn.commit()
                self.model.item(row, COL_CHK).setData(str(cursor.lastrowid), Qt.UserRole)
            else:
                cursor.execute('''UPDATE sites
                                  SET category=?,
                                      target_platform=?,
                                      site_url=?,
                                      name=?,
                                      attr=?,
                                      feature_code=?,
                                      complaint_time=?,
                                      resource_status=?,
                                      book_id=?,
                                      book_name=?,
                                      crawl_status=?
                                  WHERE id = ?''',
                               (cat, cat, url, name, attr, code, comp_time, res_status, book_id, book_name,
                                crawl_status,
                                db_id))
                conn.commit()
        except sqlite3.IntegrityError:
            pass
        finally:
            conn.close()
            self._is_updating = False

    def delete_site(self):
        selected_rows = set()
        for index in self.table.selectionModel().selectedRows():
            selected_rows.add(index.row())
        if not selected_rows: return

        reply = QMessageBox.question(self, "确认删除", f"确定要删除选中的 {len(selected_rows)} 行记录吗？",
                                     QMessageBox.Yes | QMessageBox.No)
        if reply != QMessageBox.Yes: return

        conn = get_conn()
        for proxy_row in sorted(list(selected_rows), reverse=True):
            proxy_idx = self.proxy_model.index(proxy_row, COL_CHK)
            source_idx = self.proxy_model.mapToSource(proxy_idx)
            real_row = source_idx.row()
            db_id = self.model.item(real_row, COL_CHK).data(Qt.UserRole)
            if db_id:
                conn.execute("DELETE FROM sites WHERE id=?", (db_id,))
        conn.commit()
        conn.close()
        self.refresh_table()

    def refresh_table(self):
        self._is_updating = True
        self.model.removeRows(0, self.model.rowCount())

        target_cats = self.cat_map.get(self._current_cat, ["全部资源"])

        if target_cats == ["全部资源"]:
            self.report_btn.hide()
        else:
            self.report_btn.show()

        conn = get_conn()
        cursor = conn.cursor()

        # ★ 新增 source_site 和 name，name 放最后
        query_base = ("SELECT id, target_platform, site_url, source_site, attr, feature_code, "
                      "complaint_time, resource_status, book_id, book_name, crawl_status, name FROM sites")

        if target_cats == ["全部资源"]:
            records = cursor.execute(f"{query_base} ORDER BY id ASC").fetchall()
        else:
            placeholders = ','.join(['?'] * len(target_cats))
            query = f"{query_base} WHERE (target_platform IN ({placeholders}) OR target_platform IS NULL OR target_platform = '') ORDER BY id ASC"
            params = list(target_cats)
            records = cursor.execute(query, params).fetchall()

        all_features_query = "SELECT feature_code, COUNT(*) FROM sites GROUP BY feature_code"
        feature_counts = dict(cursor.execute(all_features_query).fetchall())

        for row_data in records:
            (db_id, cat, url, publisher, attr, code, comp_time, res_status,
             b_id, b_name, crawl_status, title) = row_data

            book_display = ""
            if b_id:
                for bid, bname, _, _ in self.book_list:
                    if bid == b_id:
                        book_display = bname
                        break
            elif b_name:
                book_display = b_name

            chk_item = QStandardItem()
            chk_item.setCheckable(True)
            chk_item.setCheckState(Qt.Unchecked)
            chk_item.setData(str(db_id), Qt.UserRole)
            chk_item.setData(True, Qt.UserRole + 1)

            item_book = QStandardItem(book_display)
            item_cat = QStandardItem(cat or "")
            url_item = QStandardItem(url or "")

            # ★ 发布者只读
            item_publisher = QStandardItem(publisher or "")
            item_publisher.setEditable(False)

            item_attr = QStandardItem(attr or "")

            if feature_counts.get(code, 0) > 1:
                url_item.setText(f"{url} (重复资源链接)")
                url_item.setForeground(QColor("red"))
                url_item.setFont(QFont("Microsoft YaHei", 9, QFont.Bold))

            display_status = ""
            if not (comp_time and str(comp_time).strip()):
                display_status = "已失效" if res_status == "已失效" else ""

            item_res_status = QStandardItem(display_status)

            # ★ 标题只读
            item_title = QStandardItem(title or "")
            item_title.setEditable(False)

            self.model.appendRow([
                chk_item, item_book, item_cat, url_item, item_publisher,
                item_attr, QStandardItem(code or ""), QStandardItem(crawl_status or ""),
                QStandardItem(comp_time or ""), item_res_status, item_title,
            ])

            self.update_row_style(self.model.rowCount() - 1)

        conn.close()

        for _ in range(15):
            chk_item = QStandardItem()
            chk_item.setCheckable(True)
            chk_item.setCheckState(Qt.Unchecked)
            chk_item.setData(None, Qt.UserRole)
            self.model.appendRow([chk_item] + [QStandardItem("") for _ in range(len(TABLE_HEADERS) - 1)])

        if not getattr(self, '_initialized_filter', False):
            self._apply_filters_to_proxy()
            self._initialized_filter = True

        self.proxy_model.invalidateFilter()
        self.table.clearSelection()
        self._is_updating = False

        self._update_stats()

    def append_site_row(self, record, duplicate):
        prev = self._is_updating
        self._is_updating = True
        try:
            insert_pos = max(0, self.model.rowCount() - 15)

            chk_item = QStandardItem()
            chk_item.setCheckable(True)
            chk_item.setCheckState(Qt.Unchecked)
            chk_item.setData(str(record.get("db_id") or ""), Qt.UserRole)
            chk_item.setData(True, Qt.UserRole + 1)

            url_text = record.get("site_url") or ""
            url_item = QStandardItem(url_text)
            if duplicate:
                url_item.setText(f"{url_text} (重复资源链接)")
                url_item.setForeground(QColor("red"))
                url_item.setFont(QFont("Microsoft YaHei", 9, QFont.Bold))

            item_publisher = QStandardItem(record.get("source_site") or "")
            item_publisher.setEditable(False)

            item_title = QStandardItem(record.get("name") or "")
            item_title.setEditable(False)

            row_items = [
                chk_item,
                QStandardItem(record.get("book_name") or ""),
                QStandardItem(record.get("category") or ""),
                url_item,
                item_publisher,
                QStandardItem(record.get("attr") or ""),
                QStandardItem(record.get("feature_code") or ""),
                QStandardItem(record.get("crawl_status") or ""),
                QStandardItem(record.get("complaint_time") or ""),
                QStandardItem(record.get("resource_status") or ""),
                item_title,
            ]
            self.model.insertRow(insert_pos, row_items)

            self.update_row_style(insert_pos)
            self._is_updating = True

            self.proxy_model.invalidateFilter()
            self._update_stats()
        finally:
            self._is_updating = prev