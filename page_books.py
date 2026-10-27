import sqlite3
import os
import shutil
import uuid
import re
import csv
from PySide6.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
                               QTableWidget, QTableWidgetItem, QHeaderView, QMessageBox,
                               QAbstractItemView, QFileDialog, QSizePolicy, QDialog, QGroupBox,
                               QStyledItemDelegate, QStyle)
from PySide6.QtCore import Qt, QSettings, QUrl
from PySide6.QtGui import QDesktopServices, QColor, QPainter
from db_config import get_conn, get_current_user

# ★ 4.2：IdDelegate 泛化为 PatternCellDelegate + is_alnum_only
from ui_components import (
    CenteredCheckBoxDelegate,
    ElidedLabel,
    PatternCellDelegate,
    is_alnum_only,
    HintButton,
)
from app_paths import get_user_dir, get_user_subdir, DIR_EVIDENCES

# ★ 按钮/表格样式统一从 ui_styles 引入
from ui_styles import (
    BTN_SECONDARY,
    BTN_PRIMARY,
    BTN_PRIMARY_OUTLINE,
    BTN_DANGER_OUTLINE,
    BTN_SECONDARY_DARK,
    TABLE_BASE,
    btn_link_style,
    PRIMARY_GREEN,
    PRIMARY_GREEN_HOVER,
    TEXT_BLUE,
    TEXT_BLUE_HOVER,
    TEXT_ORANGE,
    TEXT_ORANGE_HOVER,
    DANGER_RED,
    DANGER_RED_HOVER,
    TEXT_NAVY,
    TEXT_GRAY,
    TEXT_DARK,
    BG_HOVER,
    BORDER_SOFT,
)


class EvidenceWidget(QWidget):
    def __init__(self, file_path, parent_page, row_index):
        super().__init__()
        self.file_path = file_path
        self.parent_page = parent_page
        self.row_index = row_index

        # ★ 允许宽度压缩，防止 2 个证据卡片把整个窗口撑爆
        self.setMinimumWidth(200)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)

        layout = QHBoxLayout(self)
        layout.setContentsMargins(8, 0, 8, 0)
        layout.setSpacing(8)

        if not self.file_path:
            id_item = parent_page.table.item(row_index, 1) if (parent_page and hasattr(parent_page, 'table')) else None
            db_id = id_item.data(Qt.UserRole) if id_item else ""

            if db_id:
                self.upload_btn = self._link_button("📎 上传维权书", PRIMARY_GREEN, PRIMARY_GREEN_HOVER, bold=True)
                self.upload_btn.clicked.connect(self.upload_file)
                layout.addWidget(self.upload_btn)

            layout.addStretch()
        else:
            file_name = os.path.basename(self.file_path)

            self.name_label = ElidedLabel(f"📄 {file_name}")
            self.name_label.setStyleSheet(f"color: {TEXT_NAVY}; font-size: 12px; background: transparent;")
            layout.addWidget(self.name_label, 1)

            self.preview_btn = self._link_button("预览", TEXT_BLUE, TEXT_BLUE_HOVER)
            self.preview_btn.clicked.connect(self.preview_file)

            self.reupload_btn = self._link_button("重新上传", TEXT_ORANGE, TEXT_ORANGE_HOVER)
            self.reupload_btn.clicked.connect(self.upload_file)

            self.delete_btn = self._link_button("删除", DANGER_RED, DANGER_RED_HOVER)
            self.delete_btn.clicked.connect(self.delete_file)

            layout.addWidget(self.preview_btn)
            layout.addWidget(self.reupload_btn)
            layout.addWidget(self.delete_btn)

    def _link_button(self, text, color, hover_color, bold=False):
        btn = QPushButton(text)
        btn.setCursor(Qt.PointingHandCursor)
        btn.setFlat(True)
        btn.setFocusPolicy(Qt.NoFocus)

        style = btn_link_style(color, hover_color)
        if bold:
            style += " QPushButton { font-weight: bold; }"
        btn.setStyleSheet(style)
        return btn

    def _row(self):
        table = self.parent_page.table
        idx = table.indexAt(self.mapTo(table.viewport(), self.rect().center()))
        return idx.row() if idx.isValid() else self.row_index

    def upload_file(self):
        path, _ = QFileDialog.getOpenFileName(self.window(), "选择维权书文件", "", "PDF 文件 (*.pdf)")
        if path:
            if not path.lower().endswith('.pdf'):
                QMessageBox.warning(self.window(), "格式错误", "仅支持上传 PDF 格式的维权书文件！")
                return

            row = self._row()
            item = self.parent_page.table.item(row, 1)
            db_id = item.data(Qt.UserRole) if item else None
            platform_id = self.parent_page.table.item(row, 1).text().strip() if self.parent_page.table.item(row,
                                                                                                            1) else ""

            if not db_id:
                QMessageBox.warning(self.window(), "提示", "请先填写作品名称并保存后，再上传维权书！")
                return

            try:
                user_id = get_current_user()
                if user_id:
                    data_dir = get_user_subdir(user_id, DIR_EVIDENCES)
                else:
                    base_dir = os.path.dirname(os.path.abspath(__file__))
                    data_dir = os.path.join(base_dir, "evidences_data")
                    os.makedirs(data_dir, exist_ok=True)

                prefix_id = platform_id if platform_id else f"book_{db_id}"
                safe_prefix = re.sub(r'[\\/:*?"<>|]', '_', prefix_id)

                ext = os.path.splitext(path)[1].lower()
                random_str = uuid.uuid4().hex[:4]
                new_filename = f"{safe_prefix}_作品维权书_{random_str}{ext}"
                new_file_path = os.path.join(data_dir, new_filename)

                shutil.copy2(path, new_file_path)

                conn = get_conn()
                conn.execute("UPDATE books SET evidence_path=? WHERE id=?", (new_file_path, db_id))
                conn.commit()
                conn.close()

                self._remove_old_file_if_owned(self.file_path, new_file_path)

                self.parent_page.refresh_table()
            except Exception as e:
                QMessageBox.warning(self.window(), "错误", f"保存文件失败: {str(e)}")

    def _remove_old_file_if_owned(self, old_path, new_path):
        if not old_path or not os.path.exists(old_path):
            return
        if os.path.abspath(old_path) == os.path.abspath(new_path):
            return
        conn = get_conn()
        try:
            shared = conn.execute("SELECT 1 FROM evidences WHERE file_path=?", (old_path,)).fetchone()
        finally:
            conn.close()
        if shared:
            return

        if "evidences" not in old_path.lower() and "evidences_data" not in old_path.lower():
            return
        try:
            os.remove(old_path)
        except OSError:
            pass

    def preview_file(self):
        if self.file_path and os.path.exists(self.file_path):
            QDesktopServices.openUrl(QUrl.fromLocalFile(self.file_path))
        else:
            QMessageBox.warning(self.window(), "提示", "文件不存在，可能已被移动或删除。")

    def delete_file(self):
        reply = QMessageBox.question(
            self.window(), "确认",
            "确定要删除这份维权书吗？\n\n（只删除维权书.pdf，相关原始证明图片资料不受影响）",
            QMessageBox.Yes | QMessageBox.No)
        if reply != QMessageBox.Yes:
            return

        row = self._row()
        item = self.parent_page.table.item(row, 1)
        db_id = item.data(Qt.UserRole) if item else None
        if not db_id:
            return

        conn = get_conn()
        conn.execute("UPDATE books SET evidence_path=NULL WHERE id=?", (db_id,))
        conn.commit()
        conn.close()

        self._remove_old_file_if_owned(self.file_path, new_path="")
        self.parent_page.refresh_table()


class ImportBooksDialog(QDialog):
    def __init__(self, parent_page):
        super().__init__(parent_page)
        self.parent_page = parent_page
        self.setWindowTitle("批量导入作品")
        self.setFixedSize(480, 300)
        self.setup_ui()

    def setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(15)

        group_template = QGroupBox("第一步：下载标准模板")
        group_template.setStyleSheet(f"QGroupBox {{ font-weight: bold; color: {TEXT_DARK}; }}")
        layout_template = QVBoxLayout(group_template)

        lbl_tip1 = QLabel("为了确保数据能正确导入，请先下载模板，并按照表头格式填写您的作品信息。")
        lbl_tip1.setStyleSheet(f"color: {TEXT_GRAY}; font-weight: normal;")
        lbl_tip1.setWordWrap(True)
        layout_template.addWidget(lbl_tip1)

        btn_download = QPushButton("下载 Excel 模板")
        btn_download.setStyleSheet(BTN_SECONDARY_DARK + "QPushButton { padding: 8px 12px; }")
        btn_download.clicked.connect(self.download_template)
        layout_template.addWidget(btn_download)

        layout.addWidget(group_template)

        group_upload = QGroupBox("第二步：上传已填写的文件")
        group_upload.setStyleSheet(f"QGroupBox {{ font-weight: bold; color: {TEXT_DARK}; }}")
        layout_upload = QVBoxLayout(group_upload)

        lbl_tip2 = QLabel("请选择填写好的 Excel (.xlsx) 或 CSV (.csv) 文件进行导入。如果作品名已存在，将会自动跳过。")
        lbl_tip2.setStyleSheet(f"color: {TEXT_GRAY}; font-weight: normal;")
        lbl_tip2.setWordWrap(True)
        layout_upload.addWidget(lbl_tip2)

        btn_upload = QPushButton("选择文件并导入")
        btn_upload.setStyleSheet(BTN_PRIMARY + "QPushButton { padding: 8px 12px; }")
        btn_upload.clicked.connect(self.upload_file)
        layout_upload.addWidget(btn_upload)

        layout.addWidget(group_upload)
        layout.addStretch()

    def download_template(self):
        path, _ = QFileDialog.getSaveFileName(self, "保存作品导入模板", "作品导入模板.xlsx", "Excel 文件 (*.xlsx)")
        if not path:
            return
        try:
            import openpyxl
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "作品列表"

            headers = ["平台书号(必填)", "作品名(必填)", "作品曾用名", "笔名", "权利人(必填)", "百度CP号(必填)"]
            for col, header in enumerate(headers, 1):
                ws.cell(row=1, column=col, value=header)

            for col in range(1, 7):
                ws.column_dimensions[chr(64 + col)].width = 20

            wb.save(path)
            QMessageBox.information(self, "下载成功", f"标准模板已保存至：\n{path}")
        except Exception as e:
            QMessageBox.critical(self, "失败", f"保存模板失败：{e}")

    def upload_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "选择作品数据文件", "", "表格文件 (*.csv *.xlsx)")
        if not path: return

        ext = os.path.splitext(path)[1].lower()
        rows = []
        try:
            if ext == '.csv':
                with open(path, 'r', encoding='utf-8-sig') as f:
                    rows = list(csv.reader(f))
            elif ext == '.xlsx':
                import openpyxl
                wb = openpyxl.load_workbook(path, data_only=True)
                ws = wb.active
                for row in ws.iter_rows(values_only=True):
                    if any(row):
                        rows.append([str(cell) if cell is not None else "" for cell in row])
            else:
                QMessageBox.warning(self, "提示", "不支持的文件格式！")
                return
        except Exception as e:
            QMessageBox.critical(self, "读取失败", f"读取文件出错：{e}")
            return

        if not rows or len(rows) <= 1:
            QMessageBox.warning(self, "提示", "文件为空或没有有效数据！")
            return

        header = rows[0]
        col_idx = {"platform_id": 0, "name": 1, "aliases": 2, "pen_name": 3, "rights_holder": 4, "baidu_id": 5}
        for i, h in enumerate(header):
            h_str = str(h).strip().lower()
            if "平台" in h_str or "id" in h_str and "百度" not in h_str:
                col_idx["platform_id"] = i
            elif "作品名" in h_str or "书名" in h_str:
                col_idx["name"] = i
            elif "曾用名" in h_str or "别名" in h_str:
                col_idx["aliases"] = i
            elif "笔名" in h_str:
                col_idx["pen_name"] = i
            elif "权利" in h_str:
                col_idx["rights_holder"] = i
            elif "百度" in h_str:
                col_idx["baidu_id"] = i

        conn = get_conn()
        cursor = conn.cursor()
        success_count = 0
        duplicate_count = 0
        invalid_count = 0

        for row in rows[1:]:
            if len(row) <= col_idx["name"]: continue

            name = str(row[col_idx["name"]]).strip()
            if not name or name == "None": continue

            platform_id = str(row[col_idx["platform_id"]]).strip() if len(row) > col_idx["platform_id"] else ""
            baidu_id = str(row[col_idx["baidu_id"]]).strip() if len(row) > col_idx["baidu_id"] else ""
            aliases = str(row[col_idx["aliases"]]).strip() if len(row) > col_idx["aliases"] else ""
            pen_name = str(row[col_idx["pen_name"]]).strip() if len(row) > col_idx["pen_name"] else ""
            rights = str(row[col_idx["rights_holder"]]).strip() if len(row) > col_idx["rights_holder"] else ""

            if not is_alnum_only(platform_id) or not is_alnum_only(baidu_id):
                invalid_count += 1
                continue

            try:
                cursor.execute(
                    "INSERT INTO books (platform_id, baidu_id, name, pen_name, aliases, rights_holder) VALUES (?,?,?,?,?,?)",
                    (platform_id, baidu_id, name, pen_name, aliases, rights))
                success_count += 1
            except sqlite3.IntegrityError as e:
                if "CHECK" in str(e).upper():
                    invalid_count += 1
                else:
                    duplicate_count += 1
                continue

        conn.commit()
        conn.close()

        self.parent_page.refresh_table()

        msg = f"成功导入 {success_count} 部作品！\n跳过重复 {duplicate_count} 部。"
        if invalid_count > 0:
            msg += f"\n\n⚠️ 有 {invalid_count} 条记录因【平台ID/百度CP号 含非法字符】被跳过（仅支持英文字母和数字）。"
        QMessageBox.information(self, "导入完成", msg)
        self.accept()


class BooksPage(QWidget):
    def __init__(self, main_window):
        super().__init__()
        self.main_window = main_window
        self.settings = QSettings("AuthorTools", "AntiPiracy")
        self._is_updating = False
        self.setup_ui()

    def setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(10)

        title_layout = QHBoxLayout()
        title_label = QLabel("【作品库】")
        title_label.setStyleSheet(f"font-size: 20px; font-weight: bold; color: {TEXT_DARK};")
        title_layout.addWidget(title_label)
        title_layout.addStretch()
        layout.addLayout(title_layout)

        tool_layout = QHBoxLayout()
        tool_layout.setSpacing(10)

        # ★ 加上防撑爆提示，且允许它收缩
        hint = HintButton("平台ID/百度CP号 仅支持英文字母和数字，非法会立即标红。")
        hint.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Preferred)
        tool_layout.addWidget(hint)
        tool_layout.addStretch()

        add_btn = QPushButton("➕ 新增空白行")
        add_btn.setStyleSheet(BTN_PRIMARY_OUTLINE)
        add_btn.clicked.connect(self.add_new_row)
        tool_layout.addWidget(add_btn)

        make_evidence_btn = QPushButton("制作维权书（单选）")
        make_evidence_btn.setStyleSheet(BTN_PRIMARY)
        make_evidence_btn.clicked.connect(self.go_to_make_evidence)
        tool_layout.addWidget(make_evidence_btn)

        import_btn = QPushButton("批量导入表格")
        import_btn.setStyleSheet(BTN_SECONDARY)
        import_btn.clicked.connect(self.open_import_dialog)
        tool_layout.addWidget(import_btn)

        del_btn = QPushButton("🗑️ 删除选中")
        del_btn.setStyleSheet(BTN_DANGER_OUTLINE)
        del_btn.clicked.connect(self.delete_book)
        tool_layout.addWidget(del_btn)

        layout.addLayout(tool_layout)
        layout.addSpacing(10)

        self.table = QTableWidget()
        self.table.setColumnCount(8)

        self.table.setHorizontalHeaderLabels(
            ["选择", "平台ID(必填)", "作品名(必填)", "作品曾用名", "笔名", "权利人(必填)",
             "百度CP号(必填)", "维权书(必填)"])

        # ★ 设置表头高度和列宽模式（与 links 保持一致）
        header = self.table.horizontalHeader()
        header.setFixedHeight(35)  # 表头高度改为 35

        # 模式设定：默认固定宽度，但最后一列（维权书）自动拉伸填满右侧
        header.setSectionResizeMode(QHeaderView.Interactive)
        header.setSectionResizeMode(0, QHeaderView.Fixed)
        header.setSectionResizeMode(7, QHeaderView.Stretch)

        # 初始列宽设定
        self.table.setColumnWidth(0, 40)
        self.table.setColumnWidth(1, 110)  # 平台ID
        self.table.setColumnWidth(2, 180)  # 作品名
        self.table.setColumnWidth(3, 140)  # 作品曾用名
        self.table.setColumnWidth(4, 90)  # 笔名
        self.table.setColumnWidth(5, 120)  # 权利人
        self.table.setColumnWidth(6, 140)  # 百度CP号

        self.table.verticalHeader().setDefaultSectionSize(32)  # ★ 正文行高改窄
        self.table.setAlternatingRowColors(False)  # ★ 取消正文交替灰底

        # ★ 强制覆盖表头底纹颜色，使其与 links 页面的 #F5F5F5 完全一致
        header_style = f"""
            QHeaderView::section {{
                background-color: {BG_HOVER};
                font-weight: bold;
                padding: 6px;
                border: none;
                border-bottom: 1px solid {BORDER_SOFT};
                color: {TEXT_DARK};
            }}
        """
        self.table.setStyleSheet(TABLE_BASE + header_style)

        self.table.setItemDelegateForColumn(0, CenteredCheckBoxDelegate(self.table))
        self.table.setItemDelegateForColumn(1, PatternCellDelegate(parent=self.table))
        self.table.setItemDelegateForColumn(6, PatternCellDelegate(parent=self.table))

        self.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setEditTriggers(
            QAbstractItemView.DoubleClicked | QAbstractItemView.SelectedClicked |
            QAbstractItemView.EditKeyPressed | QAbstractItemView.AnyKeyPressed)

        self.table.itemChanged.connect(self.on_item_changed)
        self.table.selectionModel().selectionChanged.connect(self.on_selection_changed)
        self.table.verticalScrollBar().valueChanged.connect(self.save_scroll)

        # ★ 允许表格被无限压窄，防止撑大主窗口
        self.table.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Expanding)

        layout.addWidget(self.table)
        self.refresh_table()

    def _row_ids_valid(self, row):
        for col in (1, 6):
            cell = self.table.item(row, col)
            if cell and not is_alnum_only(cell.text()):
                return False
        return True

    def save_scroll(self, value):
        if self._is_updating: return
        self.settings.setValue("books_scroll_pos", value)

    def on_selection_changed(self, selected, deselected):
        if self._is_updating:
            return
        self._is_updating = True
        try:
            selected_rows = set(idx.row() for idx in selected.indexes())
            deselected_rows = set(idx.row() for idx in deselected.indexes())
            deselected_rows -= selected_rows

            for r in deselected_rows:
                chk = self.table.item(r, 0)
                if chk and (chk.flags() & Qt.ItemIsUserCheckable):
                    if chk.checkState() != Qt.Unchecked:
                        chk.setCheckState(Qt.Unchecked)

            for r in selected_rows:
                chk = self.table.item(r, 0)
                if chk and (chk.flags() & Qt.ItemIsUserCheckable) and chk.data(Qt.UserRole):
                    if chk.checkState() != Qt.Checked:
                        chk.setCheckState(Qt.Checked)
        finally:
            self._is_updating = False

    def open_import_dialog(self):
        dialog = ImportBooksDialog(self)
        dialog.exec()

    def go_to_make_evidence(self):
        current_row = self.table.currentRow()
        if current_row < 0:
            QMessageBox.warning(self, "提示", "请先在表格中选中一本需要制作维权书的作品！")
            return

        id_item = self.table.item(current_row, 1)
        if not id_item:
            return

        db_id = id_item.data(Qt.UserRole)
        if not db_id:
            QMessageBox.warning(self, "提示", "请先保存该作品后再制作维权书！")
            return

        self.main_window.sidebar.setCurrentRow(5)

        if hasattr(self.main_window, 'evidence_page'):
            ep = self.main_window.evidence_page
            ep.refresh_books()
            ep.tabs.setCurrentWidget(ep.book_tab)

            idx = ep.book_tab.book_combo.findData(int(db_id))
            if idx >= 0:
                if idx == ep.book_tab.book_combo.currentIndex():
                    ep.book_tab.load_materials()
                else:
                    ep.book_tab.book_combo.setCurrentIndex(idx)

    def add_new_row(self):
        last = self.table.rowCount() - 1
        if last >= 0:
            id_item = self.table.item(last, 1)
            name_item = self.table.item(last, 2)
            if (id_item and not id_item.data(Qt.UserRole)
                    and name_item and not name_item.text().strip()):
                self.table.setCurrentCell(last, 1)
                self.table.scrollToBottom()
                return

        self.add_blank_row(set_focus=True)
        self.table.scrollToBottom()

    def _focus_row_by_db_id(self, db_id):
        target = str(db_id)
        for r in range(self.table.rowCount()):
            item = self.table.item(r, 1)
            if item and str(item.data(Qt.UserRole)) == target:
                self.table.setCurrentCell(r, 1)
                self.table.scrollToItem(item, QAbstractItemView.PositionAtCenter)
                return

    def add_blank_row(self, set_focus=False):
        prev_updating = self._is_updating
        prev_blocked = self.table.blockSignals(True)
        self._is_updating = True
        try:
            row = self.table.rowCount()
            self.table.insertRow(row)

            chk_item = QTableWidgetItem()
            chk_item.setFlags(Qt.ItemIsUserCheckable | Qt.ItemIsEnabled | Qt.ItemIsSelectable)
            chk_item.setCheckState(Qt.Unchecked)
            chk_item.setData(Qt.UserRole, "")
            self.table.setItem(row, 0, chk_item)

            id_item = QTableWidgetItem("")
            id_item.setData(Qt.UserRole, "")
            self.table.setItem(row, 1, id_item)

            for i in range(2, 7):
                self.table.setItem(row, i, QTableWidgetItem(""))

            evidence_widget = EvidenceWidget("", self, row)
            self.table.setCellWidget(row, 7, evidence_widget)

            if set_focus:
                self.table.setCurrentCell(row, 1)
        finally:
            self.table.blockSignals(prev_blocked)
            self._is_updating = prev_updating

    def delete_book(self):
        rows_to_delete = []
        for r in range(self.table.rowCount()):
            chk = self.table.item(r, 0)
            if chk and chk.checkState() == Qt.Checked:
                db_id = chk.data(Qt.UserRole)
                if db_id:
                    rows_to_delete.append((r, db_id))

        if not rows_to_delete:
            QMessageBox.information(self.window(), "提示", "请先勾选要删除的作品！（点击行即可勾选）")
            return

        reply = QMessageBox.question(
            self.window(), "确认删除",
            f"确定要删除选中的 {len(rows_to_delete)} 部作品吗？\n\n"
            f"这会同时清空这些作品在 Evidence 页的原始资料记录。",
            QMessageBox.Yes | QMessageBox.No)
        if reply == QMessageBox.Yes:
            conn = get_conn()
            for row, db_id in sorted(rows_to_delete, reverse=True):
                conn.execute("DELETE FROM books WHERE id=?", (db_id,))
                conn.execute("DELETE FROM evidences WHERE owner_type='book' AND owner_key=?", (str(db_id),))
                self.table.removeRow(row)
            conn.commit()
            conn.close()
            self.refresh_table()

    def on_item_changed(self, item):
        if self._is_updating: return
        if item.column() == 0: return
        if item.column() == 7: return

        row = item.row()

        if not self._row_ids_valid(row):
            return

        platform_id = self.table.item(row, 1).text().strip() if self.table.item(row, 1) else ""
        name = self.table.item(row, 2).text().strip() if self.table.item(row, 2) else ""

        if not name: return

        alias = self.table.item(row, 3).text().strip() if self.table.item(row, 3) else ""
        pen = self.table.item(row, 4).text().strip() if self.table.item(row, 4) else ""
        rights = self.table.item(row, 5).text().strip() if self.table.item(row, 5) else ""
        baidu_id = self.table.item(row, 6).text().strip() if self.table.item(row, 6) else ""

        id_item = self.table.item(row, 1)
        db_id = id_item.data(Qt.UserRole) if id_item else ""

        new_id = None

        conn = get_conn()
        cursor = conn.cursor()
        try:
            if not db_id:
                cursor.execute(
                    "INSERT INTO books (platform_id, baidu_id, name, pen_name, aliases, rights_holder) VALUES (?,?,?,?,?,?)",
                    (platform_id, baidu_id, name, pen, alias, rights))
                conn.commit()
                new_id = str(cursor.lastrowid)
                self.table.blockSignals(True)
                self.table.item(row, 1).setData(Qt.UserRole, new_id)
                chk = self.table.item(row, 0)
                if chk:
                    chk.setFlags(Qt.ItemIsUserCheckable | Qt.ItemIsEnabled | Qt.ItemIsSelectable)
                    chk.setData(Qt.UserRole, new_id)
                self.table.blockSignals(False)
            else:
                cursor.execute(
                    "UPDATE books SET platform_id=?, baidu_id=?, name=?, pen_name=?, aliases=?, rights_holder=? WHERE id=?",
                    (platform_id, baidu_id, name, pen, alias, rights, db_id))
                conn.commit()
        except sqlite3.IntegrityError as e:
            if "CHECK" in str(e).upper():
                QMessageBox.warning(self.window(), "格式错误", "平台ID / 百度CP号 仅支持英文字母和数字！")
                self.refresh_table()
            else:
                QMessageBox.warning(self.window(), "冲突", "作品名不能重复！")
        finally:
            conn.close()

        if not db_id:
            self.refresh_table()
            if new_id:
                self._focus_row_by_db_id(new_id)

    def refresh_table(self):
        self._is_updating = True
        self.table.blockSignals(True)
        self.table.setRowCount(0)
        conn = get_conn()

        records = conn.execute(
            "SELECT id, platform_id, baidu_id, name, pen_name, aliases, evidence_path, rights_holder FROM books").fetchall()

        for data in records:
            row = self.table.rowCount()
            self.table.insertRow(row)

            db_id = data[0]
            platform_id = data[1]
            baidu_id = data[2]
            name = data[3]
            pen = data[4]
            alias = data[5]
            evidence_path = data[6]
            rights = data[7]

            chk_item = QTableWidgetItem()
            chk_item.setFlags(Qt.ItemIsUserCheckable | Qt.ItemIsEnabled | Qt.ItemIsSelectable)
            chk_item.setCheckState(Qt.Unchecked)
            chk_item.setData(Qt.UserRole, str(db_id))
            self.table.setItem(row, 0, chk_item)

            id_item = QTableWidgetItem(str(platform_id) if platform_id else "")
            id_item.setData(Qt.UserRole, str(db_id))
            self.table.setItem(row, 1, id_item)

            self.table.setItem(row, 2, QTableWidgetItem(str(name)))
            self.table.setItem(row, 3, QTableWidgetItem(str(alias) if alias else ""))
            self.table.setItem(row, 4, QTableWidgetItem(str(pen) if pen else ""))
            self.table.setItem(row, 5, QTableWidgetItem(str(rights) if rights else ""))

            baidu_item = QTableWidgetItem(str(baidu_id) if baidu_id else "")
            self.table.setItem(row, 6, baidu_item)

            evidence_widget = EvidenceWidget(evidence_path, self, row)
            self.table.setCellWidget(row, 7, evidence_widget)

        conn.close()

        self.add_blank_row(set_focus=False)

        self.table.blockSignals(False)
        self._is_updating = False

        saved_pos = int(self.settings.value("books_scroll_pos", 0))
        self.table.verticalScrollBar().setValue(saved_pos)