# complaint_excel_utils.py
from openpyxl import Workbook
import os

from app_data import PLATFORM_BAIDU, PLATFORM_XUNLEI


def generate_platform_excel(data_list, platform_code, save_path):
    """
    根据平台代号生成对应格式的 Excel 文件
    :param data_list: 投诉数据列表 (batch)
    :param platform_code: 平台代号（PLATFORM_BAIDU / PLATFORM_XUNLEI / PLATFORM_QUARK / PLATFORM_UC）
    :param save_path: 保存的绝对路径
    """
    if not data_list:
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    if platform_code == PLATFORM_BAIDU:
        ws.append(["序号", "链接名称", "链接地址"])
        for row_idx, item in enumerate(data_list, start=2):
            ws.cell(row=row_idx, column=1, value=row_idx - 1)
            ws.cell(row=row_idx, column=2, value="盗版链接")
            ws.cell(row=row_idx, column=3, value=item["url"])

    elif platform_code == PLATFORM_XUNLEI:
        ws.append(["链接名称", "链接地址"])
        for item in data_list:
            ws.append(["盗版链接", item["url"]])

    else:
        # 夸克、UC 等默认模板
        ws.append(["侵权链接", "原创链接", "作品名（选填）"])
        for item in data_list:
            url = item["url"]
            b_name = item["book"]
            p_id = item.get("platform_id", "")
            original_url = f"https://www.jjwxc.net/onebook.php?novelid={p_id}" if p_id else "https://www.jjwxc.net/"

            ws.append([url, original_url, b_name])

    try:
        os.makedirs(os.path.dirname(os.path.abspath(save_path)), exist_ok=True)
        wb.save(save_path)
        return True
    except Exception as e:
        print(f"生成 Excel 失败: {e}")
        return False