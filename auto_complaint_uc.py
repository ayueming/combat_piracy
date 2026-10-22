import time
import re
import os
from openpyxl import Workbook
from playwright.sync_api import sync_playwright
from PySide6.QtCore import QObject, Signal

from app_paths import get_user_dir
from db_config import get_current_user


class AutoComplaintSignals(QObject):
    progress = Signal(str)


def check_and_wait_login(page, signals, stop_event, login_event, dialog_signal):
    """
    检测是否登录。UC 和夸克共用底层登录逻辑。
    """
    while True:
        if stop_event.is_set():
            return False

        try:
            page.wait_for_load_state("networkidle", timeout=5000)
        except:
            pass

        # 1. 检查是否有登录弹窗（基于你提供的DOM：div.box 或 iframe#quarkIframe）
        is_unlogged = page.locator("div.box, iframe#quarkIframe").count() > 0

        if is_unlogged:
            pass  # 明确未登录
        else:
            # 2. 没登录框，尝试等待右上角的“UC用户”出现
            try:
                page.wait_for_selector("span.user.ant-dropdown-trigger", timeout=5000)
                user_span = page.locator("span.user.ant-dropdown-trigger").first
                if "UC用户" in user_span.inner_text():
                    return True
            except Exception:
                pass

        if signals:
            signals.progress.emit("⚠️ 检测到未登录/掉线，脚本已挂起...")
        if dialog_signal:
            dialog_signal.show_login_required()

        login_event.wait()
        login_event.clear()

        if dialog_signal:
            dialog_signal.reset_progress()

        page.wait_for_timeout(2000)
        try:
            page.reload()
            page.wait_for_load_state("networkidle", timeout=10000)
        except:
            pass


def run_auto_complaint_uc(complaints_data, signals: AutoComplaintSignals = None, is_all=False,
                          stop_event=None, login_event=None, dialog_signal=None):
    def log(msg):
        print(msg)
        if signals:
            signals.progress.emit(msg)

    user_id = get_current_user()
    if not user_id:
        log("❌ 错误：未获取到当前登录用户")
        return False
    user_dir = get_user_dir(user_id)

    log("正在生成临时 Excel...")
    # 文件和夸克区分开，避免冲突
    temp_excel_path = os.path.join(user_dir, "temp", "temp_uc_import.xlsx")

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.cell(row=1, column=1, value="侵权链接")
        ws.cell(row=1, column=2, value="原创链接")
        ws.cell(row=1, column=3, value="作品名（选填）")
        start_row = 2

        for row_idx, item in enumerate(complaints_data, start=start_row):
            url = item["url"]
            book_name = item["book"]
            p_id = item.get("platform_id", "")
            original_url = f"https://www.jjwxc.net/onebook.php?novelid={p_id}" if p_id else "https://www.jjwxc.net/"

            ws.cell(row=row_idx, column=1, value=url)
            ws.cell(row=row_idx, column=2, value=original_url)
            ws.cell(row=row_idx, column=3, value=book_name)

        wb.save(temp_excel_path)
    except Exception as e:
        log(f"❌ 生成临时 Excel 失败: {e}")
        return False

    p = sync_playwright().start()
    browser = None
    try:
        # 用户数据目录区分开
        user_data_dir = os.path.join(user_dir, "login_uc_data")
        browser = p.chromium.launch_persistent_context(
            user_data_dir=user_data_dir,
            headless=False,
            args=["--start-maximized"]
        )
        page = browser.pages[0]

        log("正在打开 UC 投诉主页...")
        page.goto("https://ipp.uc.cn/#/home")

        # ★ 核心：前置登录检测
        if not check_and_wait_login(page, signals, stop_event, login_event, dialog_signal):
            log("用户中断了脚本。")
            return False

        if stop_event.is_set(): return False

        log("正在进入投诉表单...")
        try:
            btn_start = page.locator("button.ant-btn-primary:has-text('发起侵权投诉')").first
            btn_start.wait_for(state="visible", timeout=10000)
            btn_start.click()
            page.wait_for_timeout(2000)
        except Exception as e:
            log(f"❌ 找不到【发起侵权投诉】按钮: {e}")
            return False

        log("正在自动勾选固定选项...")
        try:
            page.locator("label:has-text('权利人')").click(timeout=3000)
        except:
            pass

        first_holder = complaints_data[0].get("rights_holder", "")
        if first_holder:
            log(f"正在自动选择具体权利人: {first_holder}")
            try:
                dropdown = page.locator("#obligee_id")
                if dropdown.count() > 0:
                    dropdown.scroll_into_view_if_needed()
                    page.wait_for_timeout(500)
                    dropdown.click(force=True)
                    page.wait_for_timeout(1500)

                    option = page.locator(".ant-select-dropdown-menu-item").filter(has_text=first_holder).first
                    if option.count() > 0 and option.is_visible():
                        option.click(force=True)
                    else:
                        page.keyboard.press("ArrowDown")
                        page.wait_for_timeout(500)
                        page.keyboard.press("Enter")
                page.wait_for_timeout(1000)
            except Exception as e:
                log(f"⚠️ 自动选择权利人失败: {e}")

        try:
            # UC 的投诉类型和内容类型文案与夸克一致
            page.get_by_text("知识产权").click(timeout=3000)
            page.get_by_text("著作权(含视频、图文、图集等)").click(timeout=3000)
        except:
            pass

        # ★ 核心修改：功能模块的智能匹配（忽略大小写，兼容各种叫法）
        first_cat = complaints_data[0].get("cat", "").lower()  # 统一转小写方便比对
        try:
            if "网盘" in first_cat:
                # 夸克网盘、UC网盘、uc网盘等
                page.get_by_text("UC网盘", exact=True).click(timeout=3000)
                log(f"✅ 识别到【网盘】类侵权，已选择：UC网盘")
            elif "搜索" in first_cat or "神马" in first_cat:
                # 夸克搜索、神马搜索、UC搜索等
                page.get_by_text("神马搜索", exact=True).click(timeout=3000)
                log(f"✅ 识别到【搜索】类侵权，已选择：神马搜索")
            elif "头条" in first_cat:
                # 头条内容等
                page.get_by_text("头条内容", exact=True).click(timeout=3000)
                log(f"✅ 识别到【头条】类侵权，已选择：头条内容")
            else:
                log(f"❌ 未知的平台类型: {first_cat}，请检查数据源！为防止投诉到错误模块，脚本已停止。")
                return False
        except Exception as e:
            log(f"⚠️ 选择功能模块失败: {e}")
            return False

        try:
            # 内容类型选小说
            page.get_by_text("小说", exact=True).click(timeout=3000)
        except:
            pass

        if stop_event.is_set(): return False

        log("正在执行【批量导入】...")
        try:
            page.locator("a.multi-add").first.click()
            page.wait_for_timeout(2000)

            with page.expect_file_chooser() as fc_info:
                btn_upload = page.locator("button.ant-btn-sm:has-text('上传文件'), button:has-text('上传文件')").first
                btn_upload.click()

            file_chooser = fc_info.value
            file_chooser.set_files(os.path.abspath(temp_excel_path))
            page.wait_for_timeout(4000)

            log("正在点击【解析数据】...")
            parse_btn = page.locator("button.ant-btn-primary:has-text('解析数据')").first
            parse_btn.wait_for(state="visible", timeout=5000)
            parse_btn.click(force=True)
            log("✅ 解析数据点击成功！正在等待填充...")
            page.wait_for_timeout(5000)
        except Exception as e:
            log(f"❌ 解析过程出错: {e}")
            return False

        if stop_event.is_set(): return False

        log("正在填写投诉描述...")
        try:
            first_item = complaints_data[0]
            book_name = first_item.get("book", "未知作品")
            aliases = first_item.get("aliases", "")
            p_id = first_item.get("platform_id", "")

            name_str = f"《{book_name}（原名：{aliases}）》" if aliases else f"《{book_name}》"
            url_str = f"https://www.jjwxc.net/onebook.php?novelid={p_id}" if p_id else ""

            desc_text = (
                f"晋江文学城授权作品{name_str}，"
                f"首发网址：{url_str} ，"
                f"被盗文网站非法获取，贵站出现大量盗文txt分享，持续侵犯作者的个人利益。"
                f"作者已经按照规则提交了作者的作品授权书，首发平台截图和后台截图，"
                f"希望贵站能够严肃处理这一问题，删除原帖和txt分享链接。"
            )

            # UC的placeholder与夸克一致
            desc_textarea = page.locator("textarea[placeholder*='请客观公正描述']")
            if desc_textarea.count() > 0:
                desc_textarea.fill(desc_text)
                page.wait_for_timeout(500)
        except Exception as e:
            log(f"填写描述失败: {e}")

        # ==================== 证明材料上传 ====================
        pdf_path = complaints_data[0].get("evidence_path", "")
        log(f"正在准备上传证明材料，路径: {pdf_path}")

        if pdf_path and os.path.exists(pdf_path):
            try:
                # 找到第一个“点击上传”
                upload_btn = page.locator("p", has_text="点击上传").first
                if upload_btn.count() > 0:
                    log("找到【证明文件】上传框，正在模拟点击并注入文件...")
                    with page.expect_file_chooser(timeout=5000) as fc_info:
                        upload_btn.click(force=True)
                    file_chooser = fc_info.value
                    file_chooser.set_files(os.path.abspath(pdf_path))

                    # 等待 Toast 提示出现（与你提供的截图3完全对应）
                    try:
                        page.wait_for_selector("div.ant-message:has-text('上传成功')", timeout=10000)
                        log("🎉 检测到“上传成功” Toast 提示，文件已挂载！")
                        # 等待Toast消失，防止遮挡提交按钮
                        page.wait_for_selector("div.ant-message:has-text('上传成功')", state="hidden", timeout=5000)
                    except Exception:
                        log("⚠️ 未检测到“上传成功”提示，但文件可能已塞入，继续尝试提交...")
                else:
                    log("⚠️ 未找到“点击上传”按钮！")
            except Exception as e:
                log(f"❌ PDF上传过程发生异常: {e}")
        else:
            log(f"⚠️ 本地证明材料文件不存在或为空，跳过上传。")
        # =====================================================

        if stop_event.is_set(): return False

        log("正在滚动到底部准备提交...")
        page.evaluate("window.scrollTo(0, document.body.scrollHeight);")
        page.wait_for_timeout(2000)

        log("准备点击【提交】...")
        submit_ok = False
        try:
            # 你的DOM明确显示: <span>提 交</span>，用正则匹配空格
            btn_submit = page.locator("button.ant-btn-primary").filter(has_text=re.compile(r"提\s*交")).last
            btn_submit.scroll_into_view_if_needed()
            page.wait_for_timeout(1000)

            log("正在执行真实鼠标点击...")
            btn_submit.click(force=True)

            log("✅ 已触发提交，等待响应...")
            page.wait_for_timeout(2000)

            # 二次确认弹窗（如果有）
            try:
                confirm_btn = page.locator("button.ant-btn-primary:visible").filter(
                    has_text=re.compile("确定|确认")).first
                if confirm_btn.count() > 0 and confirm_btn.is_visible():
                    confirm_btn.click(force=True)
                    log("✅ 已点击二次确认弹窗！")
                    page.wait_for_timeout(2000)
            except:
                pass

            log("正在等待提交成功的弹窗响应...")
            page.wait_for_selector("text=提交成功", timeout=15000)
            log("🎉 检测到「提交成功」弹窗，本批次投诉确认成功！")
            submit_ok = True

        except Exception as e:
            log(f"❌ 提交失败或未检测到成功提示: {e}")
            try:
                error_screenshot_dir = os.path.join(user_dir, "error_screenshots")
                os.makedirs(error_screenshot_dir, exist_ok=True)
                ts = time.strftime("%Y%m%d_%H%M%S")
                screenshot_path = os.path.join(error_screenshot_dir, f"uc_fail_{ts}.png")
                page.screenshot(path=screenshot_path, full_page=True)
                log(f"📸 已保存失败现场截图: {screenshot_path}")
            except:
                pass

        if not submit_ok:
            return False

        if is_all:
            log("✅ 提交完成！正在刷新页面回到首页，准备处理下一批...")
            page.wait_for_timeout(2000)
            page.goto("https://ipp.uc.cn/#/home")
            page.wait_for_load_state("networkidle", timeout=15000)
            return True
        else:
            log("👉 单次投诉模式，请确认结果后【手动关闭浏览器窗口】。")
            return True

    except Exception as e:
        log(f"❌ 运行遇到问题：{e}")
        return False
    finally:
        if os.path.exists(temp_excel_path):
            try:
                os.remove(temp_excel_path)
            except:
                pass
        if browser:
            if is_all and not stop_event.is_set():
                try:
                    browser.close()
                except:
                    pass
            else:
                print("（浏览器已保留，未自动关闭）")
        p.stop()