import os
import shutil
import uuid
import re
from PySide6.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
                               QComboBox, QGridLayout, QFileDialog, QMessageBox,
                               QGroupBox, QScrollArea, QFrame, QSizePolicy, QTabWidget)
from PySide6.QtCore import Qt, QUrl
from PySide6.QtGui import QDesktopServices, QPixmap, QImageReader
from db_config import get_conn

# ★ 引入路径管理
from app_paths import DATA_DIR
# ★ 引入权限管理器
from license_manager import license_mgr

HOLDER_EVIDENCE_TYPES = {
    "id_card_front": "1、身份证-人像面",
    "id_card_back": "2、身份证-国徽面",
}

BOOK_EVIDENCE_TYPES = {
    "site_backend": "1、作品独家发布网站后台页面-点我更新页",
    "site_backend_doc": "2、作品独家发布网站后台页面-文案管理页",
    "site_home": "3、作品独家发布平台首页",
    "auth_pdf": "4、作品授权书",
}

EVIDENCE_TYPES = {**HOLDER_EVIDENCE_TYPES, **BOOK_EVIDENCE_TYPES}

UNASSIGNED_HOLDER_KEY = "__default__"
UNASSIGNED_HOLDER_LABEL = "（未指定 / 旧数据）"

ALLOWED_IMAGE_EXTS = ('.png', '.jpg', '.jpeg', '.bmp')


def get_data_dir():
    data_dir = os.path.join(DATA_DIR, "evidences_data")
    os.makedirs(data_dir, exist_ok=True)
    return data_dir


def resolve_evidence_path(raw_path):
    if not raw_path:
        return ""
    if os.path.exists(raw_path):
        return raw_path

    filename = os.path.basename(raw_path)
    new_path = os.path.join(get_data_dir(), filename)
    if os.path.exists(new_path):
        try:
            conn = get_conn()
            conn.execute("UPDATE evidences SET file_path=? WHERE file_path=?", (new_path, raw_path))
            conn.execute("UPDATE books SET evidence_path=? WHERE evidence_path=?", (new_path, raw_path))
            conn.commit()
            conn.close()
        except Exception:
            pass
        return new_path
    return raw_path


# ================= 通用证据卡片 =================
class EvidenceWidget(QFrame):
    def __init__(self, ev_type, title, file_path, owner_type, owner_key, prefix, parent_page):
        super().__init__()
        self.ev_type = ev_type
        self.file_path = resolve_evidence_path(file_path)
        self.owner_type = owner_type
        self.owner_key = owner_key
        self.prefix = prefix
        self.parent_page = parent_page

        self.setFrameShape(QFrame.NoFrame)
        self.setStyleSheet("""
            QFrame { background-color: #FFFFFF; border: 1px solid #E9ECEF; border-radius: 6px; }
            QFrame:hover { border: 1px solid #D0D0D0; }
        """)
        # 稍微加高了一点，为了放下 3 个按钮
        self.setFixedSize(320, 310)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(15, 15, 15, 15)
        layout.setSpacing(10)

        lbl_title = QLabel()
        if self.ev_type == "auth_pdf":
            html_text = (f'<span style="font-weight:bold; color:#1F1F1F; font-size:13px;">{title}</span> '
                         f'<span style="color:#999; font-size:11px;">(最好带公章)</span>')
            lbl_title.setText(html_text)
        else:
            lbl_title.setText(title)
            lbl_title.setStyleSheet("font-weight: bold; color: #1F1F1F; font-size: 13px; border: none;")
        lbl_title.setWordWrap(True)
        layout.addWidget(lbl_title)

        self.img_label = QLabel()
        self.img_label.setAlignment(Qt.AlignCenter)
        self.img_label.setFixedSize(288, 140)
        self.img_label.setStyleSheet("background-color: #FAFAFA; border: 1px dashed #DCDCDC; border-radius: 4px;")
        layout.addWidget(self.img_label)

        self.lbl_file = QLabel()
        self.lbl_file.setStyleSheet("color: #666; font-size: 11px; border: none;")
        layout.addWidget(self.lbl_file)

        btn_layout = QHBoxLayout()
        btn_layout.setSpacing(6)

        # 1. 上传/更新
        self.btn_upload = QPushButton("上传/更新")
        self.btn_upload.setCursor(Qt.PointingHandCursor)
        self.btn_upload.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        self.btn_upload.clicked.connect(self.upload_file)

        # 2. 下载
        self.btn_download = QPushButton("下载")
        self.btn_download.setCursor(Qt.PointingHandCursor)
        self.btn_download.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        self.btn_download.setStyleSheet("""
            QPushButton { background-color: #F59E0B; color: white; border: none; border-radius: 4px;
                          padding: 6px 0px; font-size: 12px; font-weight: bold; }
            QPushButton:hover { background-color: #D97706; }
            QPushButton:pressed { background-color: #B45309; }
            QPushButton:disabled { background-color: #E5E7EB; color: #9CA3AF; border: none; }
        """)
        self.btn_download.clicked.connect(self.download_file)

        # 3. 查看
        self.btn_view = QPushButton("查看")
        self.btn_view.setCursor(Qt.PointingHandCursor)
        self.btn_view.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        self.btn_view.setStyleSheet("""
            QPushButton { background-color: #3B82F6; color: white; border: none; border-radius: 4px;
                          padding: 6px 0px; font-size: 12px; font-weight: bold; }
            QPushButton:hover { background-color: #2563EB; }
            QPushButton:pressed { background-color: #1D4ED8; }
            QPushButton:disabled { background-color: #E5E7EB; color: #9CA3AF; }
        """)
        self.btn_view.clicked.connect(self.view_file)

        # ★ UI 权限控制：过期禁用上传
        if license_mgr.is_expired:
            self.btn_upload.setEnabled(False)
            self.btn_upload.setStyleSheet("""
                QPushButton { background-color: #E5E7EB; color: #9CA3AF; border: none; border-radius: 4px;
                              padding: 6px 0px; font-size: 12px; font-weight: bold; }
            """)
            self.btn_upload.setToolTip("会员已过期，无法上传或更新资料。请联系管理员续费。")
        else:
            self.btn_upload.setStyleSheet("""
                QPushButton { background-color: #10B981; color: white; border: none; border-radius: 4px;
                              padding: 6px 0px; font-size: 12px; font-weight: bold; }
                QPushButton:hover { background-color: #059669; }
                QPushButton:pressed { background-color: #047857; }
            """)

        btn_layout.addWidget(self.btn_upload, 1)
        btn_layout.addWidget(self.btn_download, 1)
        btn_layout.addWidget(self.btn_view, 1)
        layout.addLayout(btn_layout)

        self.update_ui()

    def update_ui(self):
        if self.file_path and os.path.exists(self.file_path):
            file_name = os.path.basename(self.file_path)
            fm = self.lbl_file.fontMetrics()
            self.lbl_file.setText(f"✅ 已传: {fm.elidedText(file_name, Qt.ElideMiddle, self.width() - 40)}")
            self.lbl_file.setToolTip(file_name)
            self.lbl_file.setStyleSheet("color: #10B981; font-size: 11px; font-weight: bold; border: none;")
            self.btn_view.setEnabled(True)
            self.btn_download.setEnabled(True)

            ext = os.path.splitext(self.file_path)[1].lower()
            if ext in ALLOWED_IMAGE_EXTS:
                reader = QImageReader(self.file_path)
                reader.setAutoTransform(True)
                image = reader.read()
                pixmap = QPixmap.fromImage(image)

                if not pixmap.isNull():
                    self.img_label.setPixmap(pixmap.scaled(self.img_label.size(),
                                                           Qt.KeepAspectRatio, Qt.SmoothTransformation))
                else:
                    self.img_label.setText("图片加载失败")
            elif ext == '.pdf':
                self.img_label.setText("PDF 文档已上传\n(点击下方查看)")
                self.img_label.setStyleSheet(
                    "background-color: #F8F9FA; color: #666; font-weight: bold; "
                    "border: 1px solid #E9ECEF; border-radius: 4px;")
            else:
                self.img_label.setText("未知格式\n(点击下方查看)")
        else:
            self.lbl_file.setText("❌ 未上传")
            self.lbl_file.setStyleSheet("color: #E74C3C; font-size: 11px; font-weight: bold; border: none;")
            self.btn_view.setEnabled(False)
            self.btn_download.setEnabled(False)
            self.img_label.clear()
            self.img_label.setText("暂无截图")
            self.img_label.setStyleSheet(
                "background-color: #FAFAFA; color: #999; border: 1px dashed #DCDCDC; border-radius: 4px;")

    def upload_file(self):
        # ★ 逻辑层拦截
        if license_mgr.is_expired:
            QMessageBox.warning(self, "功能受限", "您的授权已过期，无法上传新的证明材料。")
            return

        path, _ = QFileDialog.getOpenFileName(
            self, f"选择 {EVIDENCE_TYPES.get(self.ev_type, '')}", "",
            "图片文件 (*.png *.jpg *.jpeg *.bmp)")
        if not path:
            return

        if not path.lower().endswith(ALLOWED_IMAGE_EXTS):
            QMessageBox.warning(self, "格式错误",
                                "维权资料仅支持上传图片格式（PNG / JPG / JPEG / BMP）！\n"
                                "如需上传 PDF 维权书，请前往【作品管理】页面操作。")
            return

        data_dir = None
        try:
            data_dir = get_data_dir()
            abs_src = os.path.abspath(path)
            abs_data_dir = os.path.abspath(data_dir)

            if abs_src.startswith(abs_data_dir):
                new_file_path = abs_src
            else:
                type_name = EVIDENCE_TYPES.get(self.ev_type, "未知资料")
                safe_prefix = re.sub(r'[\\/:*?"<>|]', '_', str(self.prefix) if self.prefix else "未命名")
                safe_type = re.sub(r'[\\/:*?"<>|]', '_', type_name)
                ext = os.path.splitext(path)[1].lower()
                random_str = uuid.uuid4().hex[:4]

                new_filename = f"{safe_prefix}_{safe_type}_{random_str}{ext}"
                new_file_path = os.path.join(data_dir, new_filename)

                try:
                    from PIL import Image, ImageOps
                    with Image.open(path) as img:
                        img = ImageOps.exif_transpose(img)
                        if img.mode in ("RGBA", "P"):
                            if new_file_path.lower().endswith('.png'):
                                img = img.convert("RGBA")
                            else:
                                img = img.convert("RGB")
                        img.save(new_file_path, optimize=True)
                except Exception:
                    shutil.copy2(path, new_file_path)

            conn = get_conn()
            conn.execute(
                "DELETE FROM evidences WHERE owner_type=? AND owner_key=? AND ev_type=?",
                (self.owner_type, self.owner_key, self.ev_type))
            conn.execute(
                "INSERT INTO evidences (owner_type, owner_key, ev_type, file_path) VALUES (?,?,?,?)",
                (self.owner_type, self.owner_key, self.ev_type, new_file_path))
            conn.commit()
            conn.close()

            if self.file_path and os.path.exists(self.file_path):
                old_abs = os.path.abspath(self.file_path)
                new_abs = os.path.abspath(new_file_path)
                if old_abs != new_abs and "evidences_data" in old_abs:
                    conn2 = get_conn()
                    try:
                        shared = conn2.execute(
                            "SELECT 1 FROM evidences WHERE file_path=?", (self.file_path,)
                        ).fetchone()
                    finally:
                        conn2.close()
                    if not shared:
                        try:
                            os.remove(self.file_path)
                        except OSError:
                            pass

            self.file_path = new_file_path
            self.update_ui()
            self.parent_page.refresh_current()
        except Exception as e:
            QMessageBox.warning(self, "错误", f"保存文件失败: {e}\n\n目标目录: {data_dir}")

    def download_file(self):
        if not self.file_path or not os.path.exists(self.file_path):
            return

        file_name = os.path.basename(self.file_path)
        save_path, _ = QFileDialog.getSaveFileName(self, "下载资料到...", file_name, "所有文件 (*.*)")
        if save_path:
            try:
                shutil.copy2(self.file_path, save_path)
                QMessageBox.information(self, "下载成功", f"文件已成功保存至：\n{save_path}")
            except Exception as e:
                QMessageBox.critical(self, "下载失败", f"保存文件时出错：\n{e}")

    def view_file(self):
        if self.file_path and os.path.exists(self.file_path):
            QDesktopServices.openUrl(QUrl.fromLocalFile(self.file_path))


# ================= Tab 1: 权利人材料 =================
class HolderTab(QWidget):
    def __init__(self, parent_page):
        super().__init__()
        self.parent_page = parent_page
        self.setup_ui()

    def setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(15)

        top_layout = QHBoxLayout()
        lbl = QLabel("选择权利人:")
        lbl.setStyleSheet("font-weight: bold; color: #333; font-size: 15px;")
        top_layout.addWidget(lbl)

        self.holder_combo = QComboBox()
        self.holder_combo.setMinimumWidth(250)
        self.holder_combo.setStyleSheet("""
            QComboBox { border: 1px solid #CCCCCC; border-radius: 4px; padding: 6px 10px;
                        background: #FFF; color: #333; font-size: 14px; }
            QComboBox:hover { border: 1px solid #1F1F1F; }
            QComboBox QAbstractItemView { border: 1px solid #CCCCCC; background: #FFF;
                                          selection-background-color: #F0F0F0; selection-color: #000; }
        """)
        self.holder_combo.currentIndexChanged.connect(self.load_materials)
        top_layout.addWidget(self.holder_combo)
        top_layout.addStretch()
        layout.addLayout(top_layout)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setStyleSheet("QScrollArea { border: none; background-color: #FFFFFF; }")
        self.content_widget = QWidget()
        self.content_widget.setStyleSheet("background-color: #FFFFFF;")
        self.content_layout = QVBoxLayout(self.content_widget)
        scroll.setWidget(self.content_widget)
        layout.addWidget(scroll)

    def refresh_holders(self):
        self.holder_combo.blockSignals(True)
        current = self.holder_combo.currentData()
        self.holder_combo.clear()

        conn = get_conn()
        holders = [r[0] for r in conn.execute(
            "SELECT DISTINCT rights_holder FROM books "
            "WHERE rights_holder IS NOT NULL AND rights_holder != '' ORDER BY rights_holder"
        ).fetchall()]
        has_default = conn.execute(
            "SELECT 1 FROM evidences WHERE owner_type='holder' AND owner_key=? LIMIT 1",
            (UNASSIGNED_HOLDER_KEY,)
        ).fetchone()
        conn.close()

        for h in holders:
            self.holder_combo.addItem(h, h)
        if has_default:
            self.holder_combo.addItem(UNASSIGNED_HOLDER_LABEL, UNASSIGNED_HOLDER_KEY)
        if self.holder_combo.count() == 0:
            self.holder_combo.addItem("（暂无权利人，请先在【作品管理】填写）", "")

        if current:
            idx = self.holder_combo.findData(current)
            if idx >= 0:
                self.holder_combo.setCurrentIndex(idx)
        self.holder_combo.blockSignals(False)

    def load_materials(self):
        while self.content_layout.count():
            item = self.content_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()

        holder = self.holder_combo.currentData()
        if not holder:
            lbl = QLabel("暂无可用于维权的权利人信息。\n请先在【作品管理】中为作品填写权利人名称。")
            lbl.setAlignment(Qt.AlignCenter)
            lbl.setStyleSheet("color: #999; padding: 30px; font-size: 14px;")
            self.content_layout.addWidget(lbl)
            self.content_layout.addStretch()
            return

        conn = get_conn()
        records = dict(conn.execute(
            "SELECT ev_type, file_path FROM evidences WHERE owner_type='holder' AND owner_key=?",
            (holder,)
        ).fetchall())
        conn.close()

        display_name = self.holder_combo.currentText()
        group = QGroupBox(f"权利人身份证明材料: {display_name}")
        group.setStyleSheet("""
            QGroupBox { font-weight: bold; font-size: 14px; color: #1F1F1F;
                        border: 1px solid #E9ECEF; border-radius: 6px;
                        margin-top: 10px; background-color: #F8F9FA; }
            QGroupBox::title { subcontrol-origin: margin; subcontrol-position: top left;
                               padding: 0 5px; left: 15px;
                               color: #1F1F1F; background-color: #FFFFFF; }
        """)
        layout_g = QHBoxLayout(group)
        layout_g.setContentsMargins(15, 20, 15, 15)
        layout_g.setSpacing(15)

        prefix = display_name if holder != UNASSIGNED_HOLDER_KEY else "未指定"
        for k in ["id_card_front", "id_card_back"]:
            path = records.get(k, "")
            w = EvidenceWidget(k, HOLDER_EVIDENCE_TYPES[k], path, "holder", holder, prefix, self)
            layout_g.addWidget(w)
        layout_g.addStretch()

        self.content_layout.addWidget(group)
        self.content_layout.addStretch()

    def refresh_current(self):
        self.load_materials()


# ================= Tab 2: 书籍材料 =================
class BookTab(QWidget):
    def __init__(self, parent_page):
        super().__init__()
        self.parent_page = parent_page
        self.setup_ui()

    def setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(15)

        top_layout = QHBoxLayout()
        lbl = QLabel("选择作品:")
        lbl.setStyleSheet("font-weight: bold; color: #333; font-size: 15px;")
        top_layout.addWidget(lbl)

        self.book_combo = QComboBox()
        self.book_combo.setMinimumWidth(250)
        self.book_combo.setStyleSheet("""
            QComboBox { border: 1px solid #CCCCCC; border-radius: 4px; padding: 6px 10px;
                        background: #FFF; color: #333; font-size: 14px; }
            QComboBox:hover { border: 1px solid #1F1F1F; }
            QComboBox QAbstractItemView { border: 1px solid #CCCCCC; background: #FFF;
                                          selection-background-color: #F0F0F0; selection-color: #000; }
        """)
        self.book_combo.currentIndexChanged.connect(self.load_materials)
        top_layout.addWidget(self.book_combo)
        top_layout.addStretch()

        self.view_weiquan_btn = QPushButton("预览维权书")
        self.view_weiquan_btn.setStyleSheet(
            "QPushButton { background-color: #10B981; color: white; padding: 6px 16px; "
            "border-radius: 4px; font-weight: bold; border: none; } "
            "QPushButton:hover { background-color: #059669; }")
        self.view_weiquan_btn.clicked.connect(self.open_weiquan_pdf)
        top_layout.addWidget(self.view_weiquan_btn)

        self.export_pdf_btn = QPushButton("保存为维权书")
        self.export_pdf_btn.clicked.connect(self.save_to_weiquan_pdf)

        self.export_docx_btn = QPushButton("导出维权书Word")
        self.export_docx_btn.clicked.connect(self.export_word)

        # ★ UI 层拦截：如果过期，禁用生成按钮
        if license_mgr.is_expired:
            expired_style = """
                QPushButton { background-color: #E5E7EB; color: #9CA3AF; 
                padding: 6px 16px; border-radius: 4px; font-weight: bold; border: none; }
            """
            self.export_pdf_btn.setStyleSheet(expired_style)
            self.export_pdf_btn.setEnabled(False)
            self.export_pdf_btn.setToolTip("会员已过期，无法生成维权书。请联系管理员续费。")

            self.export_docx_btn.setStyleSheet(expired_style)
            self.export_docx_btn.setEnabled(False)
            self.export_docx_btn.setToolTip("会员已过期，无法导出 Word。请联系管理员续费。")
        else:
            self.export_pdf_btn.setStyleSheet(
                "QPushButton { background-color: #c0392b; color: white; padding: 6px 16px; "
                "border-radius: 4px; font-weight: bold; border: none; } "
                "QPushButton:hover { background-color: #962d22; }")

            self.export_docx_btn.setStyleSheet(
                "QPushButton { background-color: #3B82F6; color: white; padding: 6px 16px; "
                "border-radius: 4px; font-weight: bold; border: none; } "
                "QPushButton:hover { background-color: #2563EB; }")

        top_layout.addWidget(self.export_pdf_btn)
        top_layout.addWidget(self.export_docx_btn)

        layout.addLayout(top_layout)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setStyleSheet("QScrollArea { border: none; background-color: #FFFFFF; }")
        self.content_widget = QWidget()
        self.content_widget.setStyleSheet("background-color: #FFFFFF;")
        self.content_layout = QVBoxLayout(self.content_widget)
        scroll.setWidget(self.content_widget)
        layout.addWidget(scroll)

    def refresh_books(self):
        self.book_combo.blockSignals(True)
        current_id = self.book_combo.currentData()
        self.book_combo.clear()
        conn = get_conn()
        for b_id, b_name in conn.execute("SELECT id, name FROM books ORDER BY id").fetchall():
            self.book_combo.addItem(b_name, b_id)
        conn.close()
        if self.book_combo.count() == 0:
            self.book_combo.addItem("（暂无作品）", 0)
        if current_id is not None:
            idx = self.book_combo.findData(current_id)
            if idx >= 0:
                self.book_combo.setCurrentIndex(idx)
        self.book_combo.blockSignals(False)

    def load_materials(self):
        while self.content_layout.count():
            item = self.content_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()

        book_id = self.book_combo.currentData()
        if not book_id:
            lbl = QLabel("暂无作品，请先在【作品管理】添加作品。")
            lbl.setAlignment(Qt.AlignCenter)
            lbl.setStyleSheet("color: #999; padding: 30px; font-size: 14px;")
            self.content_layout.addWidget(lbl)
            self.content_layout.addStretch()
            return

        conn = get_conn()
        res = conn.execute("SELECT platform_id FROM books WHERE id=?", (book_id,)).fetchone()
        platform_id = res[0] if res else ""
        records = dict(conn.execute(
            "SELECT ev_type, file_path FROM evidences WHERE owner_type='book' AND owner_key=?",
            (str(book_id),)
        ).fetchall())
        conn.close()

        group = QGroupBox(f"作品专属资料: {self.book_combo.currentText()}")
        group.setStyleSheet("""
            QGroupBox { font-weight: bold; font-size: 14px; color: #1F1F1F;
                        border: 1px solid #E9ECEF; border-radius: 6px;
                        margin-top: 10px; background-color: #F8F9FA; }
            QGroupBox::title { subcontrol-origin: margin; subcontrol-position: top left;
                               padding: 0 5px; left: 15px;
                               color: #1F1F1F; background-color: #FFFFFF; }
        """)
        layout_g = QGridLayout(group)
        layout_g.setSpacing(15)
        layout_g.setContentsMargins(15, 20, 15, 15)

        prefix = str(platform_id) if platform_id else f"book_{book_id}"
        for i, k in enumerate(["site_backend", "site_backend_doc", "site_home", "auth_pdf"]):
            path = records.get(k, "")
            w = EvidenceWidget(k, BOOK_EVIDENCE_TYPES[k], path, "book", str(book_id), prefix, self)
            layout_g.addWidget(w, i // 2, i % 2)

        self.content_layout.addWidget(group)
        self.content_layout.addStretch()

    def refresh_current(self):
        self.load_materials()

    def get_evidence_data(self, book_id):
        conn = get_conn()
        res = conn.execute("SELECT rights_holder FROM books WHERE id=?", (book_id,)).fetchone()
        rights_holder = res[0] if res and res[0] else ""
        holder_ev = {}
        if rights_holder:
            holder_ev = dict(conn.execute(
                "SELECT ev_type, file_path FROM evidences WHERE owner_type='holder' AND owner_key=?",
                (rights_holder,)
            ).fetchall())
        book_ev = dict(conn.execute(
            "SELECT ev_type, file_path FROM evidences WHERE owner_type='book' AND owner_key=?",
            (str(book_id),)
        ).fetchall())
        conn.close()

        for k in holder_ev:
            holder_ev[k] = resolve_evidence_path(holder_ev[k])
        for k in book_ev:
            book_ev[k] = resolve_evidence_path(book_ev[k])

        return holder_ev, book_ev

    def open_weiquan_pdf(self):
        book_id = self.book_combo.currentData()
        if not book_id:
            QMessageBox.warning(self, "提示", "请先选择一本具体的作品！")
            return
        conn = get_conn()
        res = conn.execute("SELECT evidence_path FROM books WHERE id=?", (book_id,)).fetchone()
        conn.close()

        path = resolve_evidence_path(res[0]) if res and res[0] else ""

        if path and os.path.exists(path):
            QDesktopServices.openUrl(QUrl.fromLocalFile(path))
        else:
            QMessageBox.warning(self, "提示", "该作品还未上传维权书。")

    def export_word(self):
        # ★ 逻辑层拦截
        if license_mgr.is_expired:
            QMessageBox.warning(self, "功能受限", "您的授权已过期，无法导出 Word 维权书。请联系管理员续费。")
            return

        book_id = self.book_combo.currentData()
        if not book_id:
            QMessageBox.warning(self, "提示", "请先选择一本具体的作品！")
            return

        reply = QMessageBox.question(self, "导入选项",
                                     "是否需要在维权书中包含【身份证明材料】（身份证正反面）？",
                                     QMessageBox.Yes | QMessageBox.No)
        include_id = (reply == QMessageBox.Yes)
        suffix = "_含身份证" if include_id else "_纯版权"
        default_filename = f"{self.book_combo.currentText()}_维权资料{suffix}.docx"

        path, _ = QFileDialog.getSaveFileName(self, "保存 Word", default_filename, "Word 文件 (*.docx)")
        if not path:
            return

        try:
            from docx import Document
            from docx.shared import Inches, Pt
            from docx.enum.text import WD_BREAK
            from PIL import Image, ImageOps
            import io

            doc = Document()
            holder_ev, book_ev = self.get_evidence_data(book_id)

            items = [(BOOK_EVIDENCE_TYPES[k], book_ev.get(k, "")) for k in
                     ["site_backend", "site_backend_doc", "site_home", "auth_pdf"]]
            if include_id:
                items.append((HOLDER_EVIDENCE_TYPES["id_card_front"], holder_ev.get("id_card_front", "")))
                items.append((HOLDER_EVIDENCE_TYPES["id_card_back"], holder_ev.get("id_card_back", "")))

            for idx, (title, file_path) in enumerate(items):
                p_title = doc.add_paragraph()
                p_title.paragraph_format.space_after = Pt(0)
                run = p_title.add_run(title)
                run.bold = True
                run.font.size = Pt(16)
                run.font.name = '微软雅黑'

                p_img = doc.add_paragraph()
                p_img.paragraph_format.space_before = Pt(0)
                p_img.paragraph_format.space_after = Pt(0)

                if file_path and os.path.exists(file_path) and not file_path.lower().endswith('.pdf'):
                    try:
                        with Image.open(file_path) as img:
                            img = ImageOps.exif_transpose(img)
                            if img.mode in ("RGBA", "P"):
                                img = img.convert("RGB")
                            img_byte_arr = io.BytesIO()
                            img.save(img_byte_arr, format='JPEG')
                            img_byte_arr.seek(0)

                            w_px, h_px = img.size
                        ar = w_px / h_px
                        max_w, max_h = 5.7, 8.5
                        if ar > (max_w / max_h):
                            fw, fh = Inches(max_w), Inches(max_w / ar)
                        else:
                            fh, fw = Inches(max_h), Inches(max_h * ar)
                        p_img.add_run().add_picture(img_byte_arr, width=fw, height=fh)
                    except Exception as e:
                        print(f"Word插图失败: {e}")
                        p_img.add_run("（图片解析失败，请重试）")
                else:
                    p_img.add_run("（缺失：未上传截图）")

                if idx < len(items) - 1:
                    p_img.add_run().add_break(WD_BREAK.PAGE)

            doc.save(path)
            QMessageBox.information(self, "成功", "Word 维权书已生成！")
            QDesktopServices.openUrl(QUrl.fromLocalFile(path))
        except Exception as e:
            QMessageBox.critical(self, "失败", f"导入失败：{e}")

    def save_to_weiquan_pdf(self):
        # ★ 逻辑层拦截
        if license_mgr.is_expired:
            QMessageBox.warning(self, "功能受限", "您的授权已过期，无法生成 PDF 维权书。请联系管理员续费。")
            return

        book_id = self.book_combo.currentData()
        if not book_id:
            QMessageBox.warning(self, "提示", "请先选择一本具体的作品！")
            return

        conn = get_conn()
        res = conn.execute("SELECT platform_id, evidence_path FROM books WHERE id=?", (book_id,)).fetchone()
        conn.close()
        platform_id = res[0] if res else ""
        old_evidence_path = resolve_evidence_path(res[1]) if res and res[1] else ""

        if old_evidence_path and os.path.exists(old_evidence_path):
            reply = QMessageBox.question(
                self, "覆盖确认",
                "该作品已存在维权书。\n\n继续操作将会生成新的 PDF 并覆盖原有的维权书，是否继续？",
                QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
            if reply != QMessageBox.Yes:
                return

        reply = QMessageBox.question(self, "保存选项",
                                     "是否需要在维权书中包含【身份证明材料】（身份证正反面）？",
                                     QMessageBox.Yes | QMessageBox.No)
        include_id = (reply == QMessageBox.Yes)

        data_dir = get_data_dir()
        prefix = str(platform_id) if platform_id else f"book_{book_id}"
        safe_prefix = re.sub(r'[\\/:*?"<>|]', '_', prefix)
        new_filename = f"{safe_prefix}_作品维权书_{uuid.uuid4().hex[:4]}.pdf"
        new_file_path = os.path.join(data_dir, new_filename)

        try:
            from fpdf import FPDF
            from PIL import Image, ImageOps
            import tempfile

            pdf = FPDF()
            font_loaded = False
            for fp in [r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simhei.ttf",
                       r"C:\Windows\Fonts\simsun.ttc"]:
                if os.path.exists(fp):
                    try:
                        pdf.add_font("chinese", "", fp, uni=True)
                        pdf.add_font("chinese", "B", fp, uni=True)
                        font_loaded = True
                        break
                    except Exception:
                        pass

            holder_ev, book_ev = self.get_evidence_data(book_id)
            items = [(BOOK_EVIDENCE_TYPES[k], book_ev.get(k, "")) for k in
                     ["site_backend", "site_backend_doc", "site_home", "auth_pdf"]]
            if include_id:
                items.append((HOLDER_EVIDENCE_TYPES["id_card_front"], holder_ev.get("id_card_front", "")))
                items.append((HOLDER_EVIDENCE_TYPES["id_card_back"], holder_ev.get("id_card_back", "")))

            for title, file_path in items:
                pdf.add_page()
                pdf.set_font("chinese" if font_loaded else "helvetica", size=16, style="B")
                pdf.multi_cell(0, 10, txt=title)
                pdf.ln(2)

                if file_path and os.path.exists(file_path) and not file_path.lower().endswith('.pdf'):
                    try:
                        with Image.open(file_path) as img:
                            img = ImageOps.exif_transpose(img)
                            if img.mode in ("RGBA", "P"):
                                img = img.convert("RGB")

                            with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as tmp:
                                img.save(tmp.name, format='JPEG')
                                tmp_path = tmp.name

                            w_px, h_px = img.size

                        max_w, max_h = 146, 230
                        ar = w_px / h_px
                        if ar > (max_w / max_h):
                            fw, fh = max_w, max_w / ar
                        else:
                            fh, fw = max_h, max_h * ar
                        cur_y = pdf.get_y()
                        pdf.image(tmp_path, w=fw, h=fh)
                        pdf.set_y(cur_y + fh)

                        os.remove(tmp_path)
                    except Exception as e:
                        print(f"PDF插图失败: {e}")
                else:
                    pdf.set_font("chinese" if font_loaded else "helvetica", size=12)
                    pdf.set_text_color(255, 0, 0)
                    pdf.cell(0, 10, txt="(缺失：未上传截图)", ln=True)
                    pdf.set_text_color(0, 0, 0)

            pdf.output(new_file_path)

            conn = get_conn()
            conn.execute("UPDATE books SET evidence_path=? WHERE id=?", (new_file_path, book_id))
            conn.commit()
            conn.close()

            if old_evidence_path and os.path.exists(old_evidence_path):
                old_abs = os.path.abspath(old_evidence_path)
                new_abs = os.path.abspath(new_file_path)
                if old_abs != new_abs and "evidences_data" in old_evidence_path:
                    conn2 = get_conn()
                    try:
                        shared = conn2.execute(
                            "SELECT 1 FROM evidences WHERE file_path=?", (old_evidence_path,)
                        ).fetchone()
                    finally:
                        conn2.close()
                    if not shared:
                        try:
                            os.remove(old_evidence_path)
                        except OSError:
                            pass

            QMessageBox.information(self, "保存成功", "PDF 维权书已自动合并并保存至作品库！")
            self.load_materials()

            main_win = self.window()
            if hasattr(main_win, 'books_page'):
                main_win.books_page.refresh_table()
            QDesktopServices.openUrl(QUrl.fromLocalFile(new_file_path))
        except Exception as e:
            QMessageBox.critical(self, "失败", f"保存失败：\n{e}")


# ================= 顶层容器 =================
class EvidencePage(QWidget):
    def __init__(self):
        super().__init__()
        self.setup_ui()

    def setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(0)

        self.tabs = QTabWidget()
        self.tabs.setStyleSheet("""
            QTabWidget::pane { border: 1px solid #DCDCDC; background: #FFFFFF; top: -1px; }
            QTabBar::tab { background: #F5F5F5; color: #666666; padding: 8px 24px;
                           border: 1px solid #DCDCDC; border-bottom: none;
                           border-top-left-radius: 6px; border-top-right-radius: 6px;
                           margin-right: 2px; }
            QTabBar::tab:hover { background: #EBEBEB; }
            QTabBar::tab:selected { background: #FFFFFF; color: #1F1F1F; font-weight: bold;
                                    border-bottom: 1px solid #FFFFFF; }
        """)
        layout.addWidget(self.tabs)

        self.holder_tab = HolderTab(self)
        self.book_tab = BookTab(self)
        self.tabs.addTab(self.holder_tab, "权利人材料")
        self.tabs.addTab(self.book_tab, "书籍材料")

        self.refresh_books()
        self.load_evidences()

    def refresh_books(self):
        self.holder_tab.refresh_holders()
        self.book_tab.refresh_books()

    def load_evidences(self):
        self.holder_tab.load_materials()
        self.book_tab.load_materials()